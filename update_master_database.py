import openpyxl
import os
import sys
import json
import re
import csv

sys.path.append(r'C:\Users\jiten\Desktop\class11\hlb_detail\scratch')
from krutidev2unicode import kru2uni

CENSUS_FILE = r'C:\Users\jiten\Desktop\census2026\New folder\Bhinai Main Census 2026 Revised after clean.xlsx'
PRINCIPAL_FILE = 'school principal bhinai.xlsx'
PRIVATE_SCHOOLS_FILE = r'C:\Users\jiten\Desktop\gram ganna\scraped_data\scraped_private_schools.csv'

KNOWN_REPLACEMENTS = {
    'miizkpk;Z': 'उपप्राचार्य',
    'miizkpk;Z ': 'उपप्राचार्य',
    'ofj"B v/;kid': 'वरिष्ठ अध्यापक',
    'ofjB v/;kid': 'वरिष्ठ अध्यापक',
    'ofj"B': 'वरिष्ठ',
    'ofjB': 'वरिष्ठ',
    'v/;kid': 'अध्यापक',
    "'kkjhfjd-f'k{kd": 'शारीरिक शिक्षक',
    'kkjhfjd-f\'k{kd': 'शारीरिक शिक्षक',
    'MkW-': 'डॉ.',
    'MkW': 'डॉ.',
    'Jh ': 'श्री ',
    'Jh': 'श्री',
    'Jherh ': 'श्रीमती ',
    'Jherh': 'श्रीमती',
    'jkmekfo': 'राउमावि',
    'jkmkfo': 'राउप्रावि',
    'jkizkfo': 'राप्रावि',
    'jkmizkfo': 'राउप्रावि',
    'jkmikfo': 'रामावि',
    'cMxkao': 'बडगांव',
    '¼lwj[k.M½': '(सूरखण्ड)',
    '¼': '(',
    '½': ')'
}

def clean_key(s):
    if not s:
        return ''
    return re.sub(r'[^a-z0-9]', '', str(s).lower())

def clean_and_convert(text):
    if not text:
        return ''
    s = str(text).strip()
    if not s or s.lower() == 'none' or s == '-':
        return ''

    for k, v in KNOWN_REPLACEMENTS.items():
        if k in s:
            s = s.replace(k, v)

    has_ascii = bool(re.search(r'[a-zA-Z]', s))
    if not has_ascii:
        return s.replace('¼', '(').replace('½', ')').strip()

    if s.startswith(('RJ', 'STF', 'SBIN', 'BARB', 'IFSC', 'http')) or '@' in s:
        return s

    # Leave pure English abbreviations intact if not KrutiDev
    if any(s.startswith(p) for p in ['GPS', 'GUPS', 'GSSS', 'GGSSS', 'MGGS', 'PM SHRI']):
        return s

    try:
        converted = kru2uni(s)
        for k, v in KNOWN_REPLACEMENTS.items():
            if k in converted:
                converted = converted.replace(k, v)
        converted = converted.replace('¼', '(').replace('½', ')').replace('  ', ' ').strip()
        return converted
    except Exception:
        return s

def standardize_post(raw_post):
    if not raw_post:
        return 'अध्यापक / शिक्षक'
    p = clean_and_convert(raw_post).strip()
    p_lower = p.lower()

    if 'उपप्राचार्य' in p or 'उप प्रधानाचार्य' in p or 'उपप्रधानाचार्य' in p or 'vice principal' in p_lower:
        return 'उपप्रधानाचार्य'
    if 'प्रधानाचार्य' in p or 'principal' in p_lower or 'हेडमास्टर' in p:
        return 'प्रधानाचार्य'
    if 'व्याख्याता' in p or 'प्राध्यापक' in p or 'lecturer' in p_lower:
        return p if ('व्याख्याता' in p or 'प्राध्यापक' in p) else f"व्याख्याता ({p})"
    if 'वरिष्ठ अध्यापक' in p or '2nd grade' in p_lower or 'द्वितीय श्रेणी' in p:
        return p if 'वरिष्ठ अध्यापक' in p else f"वरिष्ठ अध्यापक ({p})"
    if 'लेवल 2' in p or 'लेवल-2' in p or 'level 2' in p_lower or 'level-2' in p_lower or 'l-2' in p_lower:
        return 'अध्यापक लेवल-2'
    if 'लेवल 1' in p or 'लेवल-1' in p or 'level 1' in p_lower or 'level-1' in p_lower or 'l-1' in p_lower or 'लेवल प्रथम' in p:
        return 'अध्यापक लेवल-1'
    if 'शारीरिक' in p or 'pti' in p_lower:
        return 'शारीरिक शिक्षक (PTI)'
    if 'कंप्यूटर' in p or 'कम्प्यूटर' in p or 'computer' in p_lower:
        return 'बेसिक कंप्यूटर अनुदेशक'
    if 'प्रबोधक' in p:
        return 'प्रबोधक'
    if 'पुस्तकालयाध्यक्ष' in p or 'librarian' in p_lower:
        return 'पुस्तकालयाध्यक्ष'
    if 'प्रयोगशाला' in p or 'lab' in p_lower:
        return 'प्रयोगशाला सहायक'
    if 'वरिष्ठ लिपिक' in p or 'वरिष्ठ सहायक' in p or 'udc' in p_lower:
        return 'वरिष्ठ सहायक'
    if 'कनिष्ठ' in p or 'लिपिक' in p or 'ldc' in p_lower:
        return 'कनिष्ठ सहायक'
    if 'पंचायत शिक्षक' in p:
        return 'पंचायत शिक्षक'
    if 'विद्यालय सहायक' in p:
        return 'विद्यालय सहायक'
    if 'विशेष अध्यापक' in p or 'विशेष शिक्षक' in p:
        return 'विशेष शिक्षक'
    if 'सहायक कर्मचारी' in p or 'चतुर्थ श्रेणी' in p or 'peon' in p_lower:
        return 'सहायक कर्मचारी / चतुर्थ श्रेणी'
    if 'अध्यापक' in p:
        return 'अध्यापक'

    return p

def main():
    print("Step 1: Reading Excel Master and Mapping PEEOs...")
    wb_p = openpyxl.load_workbook(PRINCIPAL_FILE, data_only=True)
    ws3 = wb_p['Sheet3']
    ws5 = wb_p['Sheet5']
    ws6 = wb_p['Sheet6']

    # 1. Map Sheet 5: Nodal Senior Secondary Schools & Principals
    sheet5_schools = {}
    for r in range(4, ws5.max_row + 1):
        sch = ws5.cell(r, 2).value
        pname = ws5.cell(r, 3).value
        mob = ws5.cell(r, 4).value
        if sch:
            m = re.search(r'\((\d+)\)', str(sch))
            code = m.group(1) if m else ''
            sheet5_schools[code] = {
                'school_name': str(sch).strip(),
                'principal': str(pname).strip() if pname and str(pname).strip() != '--Select Teacher--' else 'प्रभारी प्रधानाचार्य',
                'mobile': str(mob).strip() if mob else '',
                'shala_darpan_code': code,
                'category': 'Govt. Senior Secondary',
                'type': 'Government'
            }

    # 2. Extract Subordinate Schools from Sheet 3
    subordinate_schools_by_peeo = {}
    for r in range(36, ws3.max_row + 1):
        peeo = ws3.cell(r, 2).value
        gp = ws3.cell(r, 4).value
        village = ws3.cell(r, 5).value
        school = ws3.cell(r, 6).value
        dise = ws3.cell(r, 7).value
        scode = ws3.cell(r, 8).value
        if peeo:
            p_clean = str(peeo).strip()
            if p_clean not in subordinate_schools_by_peeo:
                subordinate_schools_by_peeo[p_clean] = []
            subordinate_schools_by_peeo[p_clean].append({
                'school_name': str(school).strip() if school else '',
                'category': 'Govt. Elementary / Sec',
                'panchayat': str(gp).strip() if gp else '',
                'village': str(village).strip() if village else '',
                'dise_code': str(dise).strip() if dise else '',
                'shala_darpan_code': str(scode).strip() if scode else '',
                'type': 'Government'
            })

    # Emails
    emails = [str(ws6.cell(r, 1).value).strip() for r in range(1, ws6.max_row + 1) if ws6.cell(r, 1).value]

    peeo_sd_codes = {
        'PEEO BANDANWARA': '221769',
        'PEEO BARGAON': '221764',
        'PEEO BARLI': '221755',
        'PEEO BHINAY': '221780',
        'PEEO BOOBKIYA': '221763',
        'PEEO CHAPANERI': '221758',
        'PEEO CHHACHHUNDRA': '221787',
        'PEEO DEOLIYA KALAN': '221754',
        'PEEO DEVPURA': '488941',
        'PEEO DHANTOL': '221783',
        'PEEO EKALSEENGA': '221786',
        'PEEO GURHA KHURD': '221762',
        'PEEO KANAI KALAN': '221765',
        'PEEO KARATI': '221773',
        'PEEO KEROT': '221767',
        'PEEO KUMHARIYA': '221777',
        'PEEO LAMGARA': '221759',
        'PEEO NAGOLA': '221772',
        'PEEO NANDSI': '221756',
        'PEEO PADALIYA': '221766',
        'PEEO PADANGA': '221788',
        'PEEO RAMMALIA': '221785',
        'PEEO RATAKOT': '221775',
        'PEEO SINGAWAL': '221781',
        'PEEO SOBRI': '221782'
    }

    aliases = {
        'BARGAON': ['badgaon', 'bargaon'],
        'BARLI': ['badli', 'barli'],
        'BHINAY': ['bhinai', 'bhinay'],
        'DEVPURA': ['devriya', 'devpura'],
        'DEOLIYA KALAN': ['deoliya', 'deolia'],
        'EKALSEENGA': ['ekalsingha', 'ekalseenga'],
        'GURHA KHURD': ['gudha khurd', 'gurha khurd', 'bagrai'],
        'KANAI KALAN': ['kanai kala', 'kanai kalan'],
        'KARATI': ['karanti', 'karati', 'khedi'],
        'KEROT': ['kairot', 'kerot'],
        'PADALIYA': ['padliya', 'padaliya'],
        'RAMMALIA': ['rammaliya', 'rammalia'],
        'CHAPANERI': ['champaneri', 'chapaneri']
    }

    peeo_accounts = []
    sorted_peeo_keys = sorted(peeo_sd_codes.keys())

    for idx, p in enumerate(sorted_peeo_keys, 1):
        raw_name = p.replace('PEEO', '').strip()
        sd_code = peeo_sd_codes[p]
        nodal_info = sheet5_schools.get(sd_code)

        # Matched email
        me = None
        for em in emails:
            if clean_key(raw_name) in clean_key(em):
                me = em
                break
        if not me:
            me = f"{clean_key(raw_name)}school@gmail.com"

        principal_name = nodal_info['principal'] if nodal_info else 'प्रभारी प्रधानाचार्य'
        mobile = nodal_info['mobile'] if nodal_info else ''

        if 'BHINAY' in p:
            principal_name = 'AJAY KUMAR DHABAI'
            mobile = '9549240545'
            me = 'principalbhinai123@gmail.com'
        elif 'DEOLIYA' in p:
            principal_name = 'PURNIMA'
            mobile = '9414343109'
            me = 'deoliakalan105@gmail.com'

        # Build school list:
        # SCHOOL #1 MUST BE THE PEEO'S OWN NODAL SENIOR SECONDARY SCHOOL!
        schools_list = []
        if nodal_info:
            schools_list.append({
                'school_name': nodal_info['school_name'],
                'category': 'Govt. Senior Secondary (PEEO Nodal HQ)',
                'panchayat': raw_name,
                'village': raw_name,
                'dise_code': f"0821070{idx:02d}01",
                'shala_darpan_code': sd_code,
                'type': 'Government',
                'is_peeo_nodal': True
            })

        # Add subordinate schools from Sheet 3
        sub_schools = subordinate_schools_by_peeo.get(p, [])
        for sub in sub_schools:
            # Avoid duplicate if same code
            if sub['shala_darpan_code'] != sd_code:
                sub['is_peeo_nodal'] = False
                schools_list.append(sub)

        peeo_accounts.append({
            's_no': idx,
            'peeo_id': f"PEEO{idx:02d}",
            'peeo_name': p,
            'shala_darpan_code': sd_code,
            'panchayat_name': raw_name,
            'principal_incharge': principal_name,
            'mobile': mobile,
            'email': me,
            'username': sd_code,
            'alias_username': f"peeo_{clean_key(raw_name)}",
            'password': f"{clean_key(raw_name)}#2026",
            'schools': schools_list
        })

    # Step 2: Add all 46 Private Schools mapped to their respective PEEO
    print("Step 2: Adding Private Schools with contacts and PSP Codes...")
    if os.path.exists(PRIVATE_SCHOOLS_FILE):
        with open(PRIVATE_SCHOOLS_FILE, 'r', encoding='utf-8', errors='ignore') as fp:
            reader = csv.reader(fp)
            header = next(reader, None)
            for row in reader:
                if len(row) >= 4:
                    cat = row[0].strip()
                    sch_name = row[1].strip()
                    vil = row[2].strip()
                    panchayat = row[3].strip()

                    psp_m = re.search(r'\(([A-Z0-9]+)\)', sch_name)
                    psp_code = psp_m.group(1) if psp_m else ''

                    # Match to PEEO
                    matched_p = peeo_accounts[3] # default Bhinai
                    ck_panch = clean_key(panchayat)
                    ck_vil = clean_key(vil)

                    for pa in peeo_accounts:
                        ck_peeo = clean_key(pa['panchayat_name'])
                        if ck_panch in ck_peeo or ck_peeo in ck_panch or ck_vil in ck_peeo:
                            matched_p = pa
                            break

                    matched_p['schools'].append({
                        'school_name': sch_name,
                        'category': f"Private ({cat})",
                        'panchayat': panchayat,
                        'village': vil,
                        'dise_code': '',
                        'shala_darpan_code': psp_code,
                        'type': 'Private',
                        'is_peeo_nodal': False
                    })

    for pa in peeo_accounts:
        pa['school_count'] = len(pa['schools'])

    total_schools = sum(len(pa['schools']) for pa in peeo_accounts)
    print(f"Total PEEOs: {len(peeo_accounts)}, Total Schools: {total_schools}")

    # Step 3: Load and map 1048 Karmiks from Census
    print("Step 3: Loading 1048 Karmiks and mapping to schools and PEEOs...")
    wb_cen = openpyxl.load_workbook(CENSUS_FILE, data_only=True)
    ws_sup = wb_cen['Superviser']
    ws_prag = wb_cen['Praganak']
    ws1 = wb_cen['Sheet1']

    # Preload Unicode names from Supervisor and Praganak
    sso_lookup = {}
    for r in range(1, ws_sup.max_row + 1):
        name = ws_sup.cell(r, 2).value
        post = ws_sup.cell(r, 5).value
        sso = ws_sup.cell(r, 7).value
        if sso and name:
            sso_lookup[str(sso).strip()] = {
                'name': str(name).strip(),
                'post': clean_and_convert(post)
            }
    for r in range(2, ws_prag.max_row + 1):
        name = ws_prag.cell(r, 4).value
        post = ws_prag.cell(r, 7).value
        sso = ws_prag.cell(r, 9).value
        if sso and name:
            sso_lookup[str(sso).strip()] = {
                'name': str(name).strip(),
                'post': clean_and_convert(post)
            }

    # Pre-build School & PEEO Fast Matching Index
    # Map cleaned school string to (peeo_name, official_school_name)
    school_matcher = {}
    for pa in peeo_accounts:
        p_name = pa['peeo_name']
        for sch in pa['schools']:
            official_name = sch['school_name']
            sd_code = sch['shala_darpan_code']

            # Index by exact clean keys
            school_matcher[clean_key(official_name)] = (p_name, official_name)
            if sd_code:
                school_matcher[clean_key(sd_code)] = (p_name, official_name)

            # Strip parenthesis e.g. (221754)
            core_sch = re.sub(r'\(.*?\)', '', official_name).strip()
            school_matcher[clean_key(core_sch)] = (p_name, official_name)

    # Special well-known mappings from census school strings to official schools
    special_school_aliases = {
        'bhinai': ('PEEO BHINAY', 'GOVT. SENIOR SECONDARY SCHOOL BHINAI (221780)'),
        'gsssbhinai': ('PEEO BHINAY', 'GOVT. SENIOR SECONDARY SCHOOL BHINAI (221780)'),
        'mggsbhinai': ('PEEO BHINAY', 'MAHATMA GANDHI GOVT. SCHOOL BHINAI (221778)'),
        'deoliyakalan': ('PEEO DEOLIYA KALAN', 'PM SHRI GOVT. SENIOR SECONDARY SCHOOL DEOLIYA KALAN (221754)'),
        'mggsdeoliyakalan': ('PEEO DEOLIYA KALAN', 'MAHATMA GANDHI GOVT. SCHOOL DEOLIYA KALAN (221753)'),
        'badgaon': ('PEEO BARGAON', 'GOVT. SENIOR SECONDARY SCHOOL BADGAON - SURKHAND (221764)'),
        'badli': ('PEEO BARLI', 'GOVT. SENIOR SECONDARY SCHOOL BADLI (221755)'),
        'gsssbadli': ('PEEO BARLI', 'GOVT. SENIOR SECONDARY SCHOOL BADLI (221755)'),
        'bandanwara': ('PEEO BANDANWARA', 'GOVT. SENIOR SECONDARY SCHOOL BANDANWARA (221769)'),
        'mggsbandanwara': ('PEEO BANDANWARA', 'MAHATMA GANDHI GOVT. SCHOOL BANDANWARA (221770)'),
        'nandsi': ('PEEO NANDSI', 'GOVT. SENIOR SECONDARY SCHOOL NANDSI (221756)'),
        'ggsssnandsi': ('PEEO NANDSI', 'GOVT. GIRLS SENIOR SECONDARY SCHOOL NANDSI (410632)'),
        'boobkiya': ('PEEO BOOBKIYA', 'GOVT. SENIOR SECONDARY SCHOOL BOOBKIYA (221763)'),
        'chapaneri': ('PEEO CHAPANERI', 'GOVT. SENIOR SECONDARY SCHOOL CHAPANERI (221758)'),
        'chhachhundra': ('PEEO CHHACHHUNDRA', 'GOVT. SENIOR SECONDARY SCHOOL CHHACHHUNDRA (221787)'),
        'devpura': ('PEEO DEVPURA', 'GOVT. SENIOR SECONDARY SCHOOL DEVRIYA (488941)'),
        'dhantol': ('PEEO DHANTOL', 'GOVT. SENIOR SECONDARY SCHOOL DHANTOL (221783)'),
        'ekalseenga': ('PEEO EKALSEENGA', 'GOVT. SENIOR SECONDARY SCHOOL EKALSINGHA (221786)'),
        'gudhakhurd': ('PEEO GURHA KHURD', 'GOVT. SENIOR SECONDARY SCHOOL GUDHA KHURD BLOCK BHINAI DIST AJMER (221762)'),
        'kanaikalan': ('PEEO KANAI KALAN', 'GOVT. SENIOR SECONDARY SCHOOL KANAI KALA (221765)'),
        'karati': ('PEEO KARATI', 'GOVT. SENIOR SECONDARY SCHOOL KARANTI (221773)'),
        'kerot': ('PEEO KEROT', 'GOVT. SENIOR SECONDARY SCHOOL KAIROT GURJAR MOHALLA JATPURA ROAD (221767)'),
        'kumhariya': ('PEEO KUMHARIYA', 'GOVT. SENIOR SECONDARY SCHOOL KUMHARIYA THASIL BHINAI DISTRICT AJMER (221777)'),
        'lamgara': ('PEEO LAMGARA', 'GOVT. SENIOR SECONDARY SCHOOL LAMGARA (221759)'),
        'nagola': ('PEEO NAGOLA', 'GOVT. SENIOR SECONDARY SCHOOL NAGOLA (221772)'),
        'padaliya': ('PEEO PADALIYA', 'GOVT. SENIOR SECONDARY SCHOOL PADLIYA (221766)'),
        'padanga': ('PEEO PADANGA', 'GOVT. SENIOR SECONDARY SCHOOL PADANGA (221788)'),
        'rammalia': ('PEEO RAMMALIA', 'GOVT. SENIOR SECONDARY SCHOOL RAMMALIYA (221785)'),
        'ratakot': ('PEEO RATAKOT', 'GOVT. SENIOR SECONDARY SCHOOL RATAKOT (221775)'),
        'singawal': ('PEEO SINGAWAL', 'GOVT. SENIOR SECONDARY SCHOOL SINGAWAL (221781)'),
        'sobri': ('PEEO SOBRI', 'GOVT. SENIOR SECONDARY SCHOOL SOBRI (221782)')
    }

    karmiks = []
    for r in range(4, ws1.max_row + 1):
        name_raw = ws1.cell(r, 2).value
        if not name_raw:
            continue

        sso = str(ws1.cell(r, 7).value or '').strip()
        if sso == 'None':
            sso = ''

        if sso and sso in sso_lookup:
            name = sso_lookup[sso]['name']
            post = sso_lookup[sso]['post'] or clean_and_convert(ws1.cell(r, 5).value)
        else:
            name = clean_and_convert(name_raw)
            post = clean_and_convert(ws1.cell(r, 5).value)

        gender = str(ws1.cell(r, 3).value or '').strip()
        dob = str(ws1.cell(r, 4).value or '').strip()
        school_raw = str(ws1.cell(r, 6).value or '').strip()
        school_converted = clean_and_convert(school_raw)

        mob = str(ws1.cell(r, 8).value or '').strip()
        email = str(ws1.cell(r, 9).value or '').strip()
        bank_name = str(ws1.cell(r, 10).value or '').strip()
        bank_acc = str(ws1.cell(r, 11).value or '').strip()
        ifsc = str(ws1.cell(r, 12).value or '').strip()
        pan = str(ws1.cell(r, 13).value or '').strip()

        if mob.endswith('.0'): mob = mob[:-2]
        if mob in ['0', 'None', '-']: mob = ''
        if email.lower() in ['none', '-']: email = ''

        # Match to PEEO and Official School
        ck_raw = clean_key(school_raw)
        ck_conv = clean_key(school_converted)

        matched_peeo = None
        matched_school = None

        # 1. Direct alias match
        for ak, (p_res, s_res) in special_school_aliases.items():
            if ak in ck_raw or ak in ck_conv:
                matched_peeo = p_res
                matched_school = s_res
                break

        # 2. Match in school_matcher
        if not matched_school:
            for k, (p_res, s_res) in school_matcher.items():
                if len(k) >= 5 and (k in ck_raw or k in ck_conv or ck_conv in k):
                    matched_peeo = p_res
                    matched_school = s_res
                    break

        # 3. Fallback matching by PEEO name
        if not matched_peeo:
            for pa in peeo_accounts:
                p_core = clean_key(pa['panchayat_name'])
                if p_core in ck_raw or p_core in ck_conv:
                    matched_peeo = pa['peeo_name']
                    # assign to PEEO nodal school
                    matched_school = pa['schools'][0]['school_name'] if pa['schools'] else pa['peeo_name']
                    break

        if not matched_peeo:
            matched_peeo = 'PEEO BHINAY'
            matched_school = school_converted if school_converted else 'GOVT. SENIOR SECONDARY SCHOOL BHINAI (221780)'

        staff_id = f"STF{1000 + len(karmiks) + 1}"
        karmiks.append({
            'staff_id': staff_id,
            'name': name,
            'gender': gender,
            'dob': dob,
            'post': standardize_post(post),
            'school_name': matched_school if matched_school else school_converted,
            'peeo_name': matched_peeo,
            'sso_id': sso,
            'mobile': mob,
            'email': email,
            'bank_name': bank_name if bank_name != 'None' else '',
            'bank_acc': bank_acc if bank_acc != 'None' else '',
            'ifsc': ifsc if ifsc != 'None' else '',
            'pan': pan if pan != 'None' else '',
            'status': 'Active'
        })

    print(f"Total karmiks processed: {len(karmiks)}")

    admins = [
        {
            'admin_id': 'ADMIN01',
            'name': 'प्रमिला रासलोत (CBEO)',
            'post': 'मुख्य ब्लॉक शिक्षा अधिकारी (CBEO)',
            'office': 'कार्यालय मुख्य ब्लॉक शिक्षा अधिकारी, भिनाय',
            'mobile': '9414000000',
            'email': 'cbeo.bhinai.ajmer@rajasthan.gov.in',
            'username': 'cbeo_admin',
            'shala_darpan_code': '8140',
            'password': 'cbeo@2026',
            'role': 'Super Admin'
        },
        {
            'admin_id': 'ADMIN02',
            'name': 'जितेन्द्र कुमार (Jitendra Kumar)',
            'post': 'तकनीकी नोडल प्रभारी एवं व्यवस्थापक (Admin)',
            'office': 'कार्यालय मुख्य ब्लॉक शिक्षा अधिकारी, भिनाय',
            'mobile': '7073800244',
            'email': 'jitendrakumar.cbeo@gmail.com',
            'username': 'jitendra_admin',
            'shala_darpan_code': 'admin_jitendra',
            'password': 'jitendra#2026',
            'role': 'Super Admin'
        }
    ]

    master_data = {
        'cbeo_info': {
            'office_name': 'कार्यालय मुख्य ब्लॉक शिक्षा अधिकारी (CBEO), भिनाय',
            'block': 'भिनाय (BHINAI)',
            'district': 'अजमेर / केकड़ी (AJMER)',
            'nic_sd_id': '8140',
            'ifms_id': '1408',
            'cbeo_officer': 'प्रमिला रासलोत',
            'admins': admins
        },
        'admins': admins,
        'peeos': peeo_accounts,
        'staff': karmiks,
        'total_staff_count': len(karmiks),
        'total_schools_count': total_schools
    }

    # Save to JSON & JS
    with open('master_cbeo_data.json', 'w', encoding='utf-8') as f:
        json.dump(master_data, f, indent=2, ensure_ascii=False)

    with open('master_cbeo_data.js', 'w', encoding='utf-8') as f:
        f.write("const MASTER_CBEO_DATA = " + json.dumps(master_data, indent=2, ensure_ascii=False) + ";\n")

    print("SUCCESS: master_cbeo_data.json & master_cbeo_data.js successfully written!")

    # Sync to Google Drive
    print("Step 4: Syncing to Google Drive...")
    try:
        from create_cbeo_drive_backend import get_services
        drive, sheets = get_services()
        with open('drive_config.json', 'r', encoding='utf-8') as f:
            cfg = json.load(f)

        # 1_CBEO_Admin_Access_Control
        sheet1_id = cfg['sheets']['admin_access']['id']
        s1_rows = [
            ['क्र.सं.', 'ID/Code', 'नाम / कार्यालय', 'भूमिका', 'प्रभारी', 'मोबाइल नंबर', 'ईमेल ID', 'Shala Darpan Code (User ID)', 'पासवर्ड', 'कुल स्कूल']
        ]
        for a in admins:
            s1_rows.append([len(s1_rows), a['admin_id'], a['office'], a['role'], a['name'], a['mobile'], a['email'], a['shala_darpan_code'], a['password'], total_schools])
        for p in peeo_accounts:
            s1_rows.append([len(s1_rows), p['peeo_id'], p['peeo_name'], 'PEEO Incharge', p['principal_incharge'], p['mobile'], p['email'], p['shala_darpan_code'], p['password'], p['school_count']])

        sheets.spreadsheets().values().clear(spreadsheetId=sheet1_id, range="Sheet1!A1:Z500").execute()
        sheets.spreadsheets().values().update(
            spreadsheetId=sheet1_id,
            range="Sheet1!A1",
            valueInputOption="RAW",
            body={"values": s1_rows}
        ).execute()

        # 2_CBEO_Staff_Directory
        sheet2_id = cfg['sheets']['staff_directory']['id']
        s2_rows = [
            ['क्र.सं.', 'Staff ID', 'कार्मिक का नाम', 'पद', 'विद्यालय का नाम', 'संबंधित PEEO', 'मोबाइल नं.', 'ईमेल ID', 'SSO ID', 'बैंक खाता संख्या', 'IFSC कोड', 'PAN नं.', 'स्थिति']
        ]
        for idx, s in enumerate(karmiks, 1):
            s2_rows.append([
                idx, s['staff_id'], s['name'], s['post'], s['school_name'], s['peeo_name'], s['mobile'], s['email'], s['sso_id'], s['bank_acc'], s['ifsc'], s['pan'], s['status']
            ])

        sheets.spreadsheets().values().clear(spreadsheetId=sheet2_id, range="Sheet1!A1:Z2500").execute()
        sheets.spreadsheets().values().update(
            spreadsheetId=sheet2_id,
            range="Sheet1!A1",
            valueInputOption="RAW",
            body={"values": s2_rows}
        ).execute()
        print(f"Successfully synced {len(karmiks)} karmiks to Google Drive!")

    except Exception as e:
        print("Note on Google Drive Sync:", e)

if __name__ == '__main__':
    main()
