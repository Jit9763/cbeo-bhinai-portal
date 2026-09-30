import os
import sys
import json
import re
import openpyxl

from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.discovery import build

SCOPES = [
    'https://www.googleapis.com/auth/spreadsheets',
    'https://www.googleapis.com/auth/drive'
]

TOKEN_FILE = r"C:\Users\jiten\Desktop\panchayat chunav\scratch\token_panchayat.json"

def get_services():
    creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
    drive = build('drive', 'v3', credentials=creds)
    sheets = build('sheets', 'v4', credentials=creds)
    return drive, sheets

def clean_key(s):
    return re.sub(r'[^a-z0-9]', '', str(s).lower())

def extract_all_cbeo_data():
    wb_main = openpyxl.load_workbook('school principal bhinai.xlsx', data_only=True)
    ws3 = wb_main['Sheet3']
    ws5 = wb_main['Sheet5']
    ws6 = wb_main['Sheet6']

    # 1. PEEO Schools Mapping
    peeo_schools = {}
    for r in range(36, ws3.max_row+1):
        peeo = ws3.cell(r, 2).value
        gp = ws3.cell(r, 4).value
        village = ws3.cell(r, 5).value
        school = ws3.cell(r, 6).value
        dise = ws3.cell(r, 7).value
        scode = ws3.cell(r, 8).value
        if peeo:
            p_clean = str(peeo).strip()
            if p_clean not in peeo_schools:
                peeo_schools[p_clean] = []
            peeo_schools[p_clean].append({
                'school_name': str(school).strip() if school else '',
                'panchayat': str(gp).strip() if gp else '',
                'village': str(village).strip() if village else '',
                'dise_code': str(dise).strip() if dise else '',
                'school_code': str(scode).strip() if scode else ''
            })

    # 2. Principals
    principals = []
    for r in range(4, ws5.max_row+1):
        sch = ws5.cell(r, 2).value
        pname = ws5.cell(r, 3).value
        mob = ws5.cell(r, 4).value
        if sch:
            principals.append({'school': str(sch).strip(), 'name': str(pname).strip(), 'mob': str(mob).strip()})

    # 3. Emails
    emails = [str(ws6.cell(r, 1).value).strip() for r in range(1, ws6.max_row+1) if ws6.cell(r, 1).value]

    # Aliases
    aliases = {
        'BARGAON': ['badgaon', 'bargaon'],
        'BARLI': ['badli', 'barli', 'barliajm'],
        'BHINAY': ['bhinai', 'bhinay'],
        'DEVPURA': ['devriya', 'devpura'],
        'DEOLIYA KALAN': ['deoliya kalan', 'deoliakalan', 'deoliya'],
        'EKALSEENGA': ['ekalsingha', 'ekalseenga'],
        'GURHA KHURD': ['gudha khurd', 'gurha khurd', 'gudhakhurd'],
        'KANAI KALAN': ['kanai kala', 'kanai kalan', 'kanaikala'],
        'KARATI': ['karanti', 'karati'],
        'KEROT': ['kairot', 'kerot'],
        'PADALIYA': ['padliya', 'padaliya'],
        'RAMMALIA': ['rammaliya', 'rammalia'],
        'CHAPANERI': ['champaneri', 'chapaneri']
    }

    # 4. Process PEEO list
    peeo_accounts = []
    sorted_peeos = sorted(peeo_schools.keys())

    for idx, p in enumerate(sorted_peeos, 1):
        raw_name = p.replace('PEEO', '').strip()
        match_keys = [raw_name]
        for k, v in aliases.items():
            if k in raw_name or raw_name in k:
                match_keys.extend(v)

        mp = None
        for k in match_keys:
            ck = clean_key(k)
            for item in principals:
                if ck in clean_key(item['school']):
                    mp = item
                    break
            if mp:
                break

        me = None
        for k in match_keys:
            ck = clean_key(k)
            for em in emails:
                if ck in clean_key(em):
                    me = em
                    break
            if me:
                break

        # Specific fixes
        if 'BHINAY' in p:
            mp = {'name': 'AJAY KUMAR DHABAI', 'mob': '9549240545'}
            me = 'principalbhinai123@gmail.com'
        elif 'DEOLIYA' in p:
            mp = {'name': 'RAKESH KUMAR BIRAWAT', 'mob': '9829835751'}
            me = 'deoliakalan105@gmail.com'

        username = f"peeo_{raw_name.lower().replace(' ', '_')}"
        password = f"{clean_key(raw_name)}#2026"

        peeo_accounts.append({
            's_no': idx,
            'peeo_id': f"PEEO{idx:02d}",
            'peeo_name': p,
            'panchayat_name': peeo_schools[p][0]['panchayat'] if peeo_schools[p] else raw_name,
            'principal_incharge': mp['name'] if mp else 'प्रभारी प्रधानाचार्य',
            'mobile': mp['mob'] if mp else '',
            'email': me if me else f"{clean_key(raw_name)}school@gmail.com",
            'username': username,
            'password': password,
            'school_count': len(peeo_schools[p]),
            'schools': peeo_schools[p]
        })

    # 5. Extract Karmik/Staff data
    all_staff = []
    staff_id_counter = 1001

    # Add Principals as staff
    for pr in principals:
        # Match to PEEO
        matched_peeo = 'PEEO BHINAY'
        for pa in peeo_accounts:
            for s in pa['schools']:
                if clean_key(s['school_name']) in clean_key(pr['school']) or clean_key(pr['school']) in clean_key(s['school_name']):
                    matched_peeo = pa['peeo_name']
                    break

        all_staff.append({
            'staff_id': f"STF{staff_id_counter}",
            'name': pr['name'],
            'post': 'प्रधानाचार्य / Headmaster',
            'mobile': pr['mob'],
            'email': '',
            'sso_id': '',
            'school_name': pr['school'],
            'peeo_name': matched_peeo,
            'bank_acc': '',
            'ifsc': '',
            'status': 'Active'
        })
        staff_id_counter += 1

    # Add OBC Karmik staff
    obc_file = r'C:\Users\jiten\Desktop\ews survey\BLO FOR OBC SURVAY\BHINAI UPDATED OBC SURVEY KARMIT DETAIL.xlsx'
    if os.path.exists(obc_file):
        wb_obc = openpyxl.load_workbook(obc_file, data_only=True)
        ws_obc = wb_obc.active
        for r in range(5, ws_obc.max_row+1):
            part_name = ws_obc.cell(r, 3).value
            ename = ws_obc.cell(r, 4).value
            post = ws_obc.cell(r, 5).value
            mob = ws_obc.cell(r, 6).value
            sso = ws_obc.cell(r, 7).value
            if ename:
                pname_str = str(part_name) if part_name else ''
                matched_peeo = 'PEEO BHINAY'
                for pa in peeo_accounts:
                    if clean_key(pa['panchayat_name']) in clean_key(pname_str):
                        matched_peeo = pa['peeo_name']
                        break

                all_staff.append({
                    'staff_id': f"STF{staff_id_counter}",
                    'name': str(ename).strip(),
                    'post': str(post).strip() if post else 'शिक्षक (Teacher)',
                    'mobile': str(mob).strip() if mob else '',
                    'email': '',
                    'sso_id': str(sso).strip() if sso else '',
                    'school_name': pname_str,
                    'peeo_name': matched_peeo,
                    'bank_acc': '',
                    'ifsc': '',
                    'status': 'Active'
                })
                staff_id_counter += 1

    # Add Census staff
    census_file = r'C:\Users\jiten\Desktop\ews survey\ACCOUNT DEATAIL\BHINAI CANSUS KARMIK MANDAY.xlsx'
    if os.path.exists(census_file):
        wb_cen = openpyxl.load_workbook(census_file, data_only=True)
        for sname in ['Supervisors', 'Enumerators']:
            if sname in wb_cen.sheetnames:
                ws_c = wb_cen[sname]
                for r in range(4, ws_c.max_row+1):
                    name = ws_c.cell(r, 3).value
                    posting = ws_c.cell(r, 4).value
                    bank = ws_c.cell(r, 5).value
                    ifsc = ws_c.cell(r, 7).value
                    acc = ws_c.cell(r, 8).value
                    pan = ws_c.cell(r, 9).value
                    if name:
                        post_str = str(posting) if posting else ''
                        matched_peeo = 'PEEO BHINAY'
                        for pa in peeo_accounts:
                            if clean_key(pa['panchayat_name']) in clean_key(post_str):
                                matched_peeo = pa['peeo_name']
                                break

                        all_staff.append({
                            'staff_id': f"STF{staff_id_counter}",
                            'name': str(name).strip(),
                            'post': f"शिक्षक / Census {sname[:-1]}",
                            'mobile': '',
                            'email': '',
                            'sso_id': '',
                            'school_name': post_str,
                            'peeo_name': matched_peeo,
                            'bank_acc': str(acc).strip() if acc else '',
                            'ifsc': str(ifsc).strip() if ifsc else '',
                            'status': 'Active'
                        })
                        staff_id_counter += 1

    return peeo_accounts, all_staff, principals, emails

def main():
    print("1. Extracting data from Excel files...")
    peeo_accounts, all_staff, principals, emails = extract_all_cbeo_data()
    print(f"Extracted {len(peeo_accounts)} PEEOs, {len(all_staff)} Staff members.")

    print("\n2. Connecting to Google Drive...")
    drive, sheets = get_services()

    # Create or find CBEO Bhinai Portal Folder
    folder_name = "CBEO Bhinai Portal"
    res = drive.files().list(q=f"name='{folder_name}' and mimeType='application/vnd.google-apps.folder' and trashed=false").execute()
    items = res.get('files', [])
    if items:
        folder_id = items[0]['id']
        print(f"Found existing Google Drive folder: {folder_name} (ID: {folder_id})")
    else:
        folder_metadata = {
            'name': folder_name,
            'mimeType': 'application/vnd.google-apps.folder'
        }
        folder = drive.files().create(body=folder_metadata, fields='id').execute()
        folder_id = folder.get('id')
        print(f"Created new Google Drive folder: {folder_name} (ID: {folder_id})")

    # Helper function to create or get sheet
    def get_or_create_sheet(sheet_title):
        q = f"name='{sheet_title}' and '{folder_id}' in parents and trashed=false"
        r = drive.files().list(q=q).execute()
        f_items = r.get('files', [])
        if f_items:
            s_id = f_items[0]['id']
            print(f"Found existing Sheet: {sheet_title} ({s_id})")
            return s_id
        else:
            meta = {
                'name': sheet_title,
                'mimeType': 'application/vnd.google-apps.spreadsheet',
                'parents': [folder_id]
            }
            sheet_obj = drive.files().create(body=meta, fields='id').execute()
            s_id = sheet_obj.get('id')
            print(f"Created new Sheet: {sheet_title} ({s_id})")
            return s_id

    # 1. Admin Access Control Sheet
    sheet1_id = get_or_create_sheet("1_CBEO_Admin_Access_Control")
    # Populate Sheet1
    s1_rows = [
        ['क्र.सं.', 'PEEO ID', 'PEEO/कार्यालय नाम', 'पंचायत', 'प्रभारी का नाम', 'मोबाइल नंबर', 'ईमेल ID', 'यूजरनेम', 'पासवर्ड', 'भूमिका', 'कुल स्कूल'],
        [1, 'ADMIN01', 'कार्यालय मुख्य ब्लॉक शिक्षा अधिकारी (CBEO) भिनाय', 'भिनाय', 'प्रमिला रासलोत (CBEO)', '9414000000', 'cbeo.bhinai.ajmer@rajasthan.gov.in', 'cbeo_admin', 'cbeo@2026', 'Super Admin', 104]
    ]
    for p in peeo_accounts:
        s1_rows.append([
            p['s_no'] + 1,
            p['peeo_id'],
            p['peeo_name'],
            p['panchayat_name'],
            p['principal_incharge'],
            p['mobile'],
            p['email'],
            p['username'],
            p['password'],
            'PEEO Incharge',
            p['school_count']
        ])
    sheets.spreadsheets().values().update(
        spreadsheetId=sheet1_id,
        range="Sheet1!A1",
        valueInputOption="RAW",
        body={"values": s1_rows}
    ).execute()
    print("Updated 1_CBEO_Admin_Access_Control successfully!")

    # 2. Staff Directory Sheet
    sheet2_id = get_or_create_sheet("2_CBEO_Staff_Directory")
    s2_rows = [
        ['क्र.सं.', 'Staff ID', 'कार्मिक का नाम', 'पद', 'मोबाइल नं.', 'ईमेल', 'SSO ID', 'विद्यालय का नाम', 'संबंधित PEEO', 'बैंक खाता संख्या', 'IFSC कोड', 'स्थिति']
    ]
    for idx, s in enumerate(all_staff, 1):
        s2_rows.append([
            idx,
            s['staff_id'],
            s['name'],
            s['post'],
            s['mobile'],
            s['email'],
            s['sso_id'],
            s['school_name'],
            s['peeo_name'],
            s['bank_acc'],
            s['ifsc'],
            s['status']
        ])
    sheets.spreadsheets().values().update(
        spreadsheetId=sheet2_id,
        range="Sheet1!A1",
        valueInputOption="RAW",
        body={"values": s2_rows}
    ).execute()
    print(f"Updated 2_CBEO_Staff_Directory with {len(s2_rows)-1} rows successfully!")

    # 3. Information Requests Sheet
    sheet3_id = get_or_create_sheet("3_CBEO_Information_Requests")
    s3_rows = [
        # Row 1: LOCKED / UNLOCKED indicators
        ['LOCKED', 'LOCKED', 'LOCKED', 'UNLOCKED', 'UNLOCKED', 'UNLOCKED', 'UNLOCKED', 'UNLOCKED', 'LOCKED'],
        # Row 2: Field Headers
        ['क्र.सं.', 'PEEO ID', 'PEEO का नाम', 'कुल स्वीकृत कक्षा कक्ष', 'मरम्मत योग्य कक्ष संख्या', 'नवीन कक्ष आवश्यकता', 'अनुमानित व्यय (रु. लाखों में)', 'विशेष अभियुक्ति / रिमार्क', 'स्थिति (Status)']
    ]
    for idx, p in enumerate(peeo_accounts, 1):
        s3_rows.append([
            idx,
            p['peeo_id'],
            p['peeo_name'],
            "", "", "", "", "", "Pending"
        ])
    sheets.spreadsheets().values().update(
        spreadsheetId=sheet3_id,
        range="Sheet1!A1",
        valueInputOption="RAW",
        body={"values": s3_rows}
    ).execute()
    print("Updated 3_CBEO_Information_Requests successfully!")

    # 4. Audit Backup Logs Sheet
    sheet4_id = get_or_create_sheet("4_CBEO_Audit_Backup_Logs")
    s4_rows = [
        ['समय व दिनांक', 'उपयोगकर्ता (User)', 'PEEO / भूमिका', 'क्रिया (Action)', 'कार्मिक/प्रपत्र ID', 'विवरण / पुराना डेटा', 'नया डेटा / टिप्पणी']
    ]
    sheets.spreadsheets().values().update(
        spreadsheetId=sheet4_id,
        range="Sheet1!A1",
        valueInputOption="RAW",
        body={"values": s4_rows}
    ).execute()
    print("Updated 4_CBEO_Audit_Backup_Logs successfully!")

    # Save local metadata & config
    config = {
        'folder_id': folder_id,
        'folder_url': f"https://drive.google.com/drive/folders/{folder_id}",
        'sheets': {
            'admin_access': {
                'id': sheet1_id,
                'name': '1_CBEO_Admin_Access_Control',
                'url': f"https://docs.google.com/spreadsheets/d/{sheet1_id}/edit"
            },
            'staff_directory': {
                'id': sheet2_id,
                'name': '2_CBEO_Staff_Directory',
                'url': f"https://docs.google.com/spreadsheets/d/{sheet2_id}/edit"
            },
            'information_requests': {
                'id': sheet3_id,
                'name': '3_CBEO_Information_Requests',
                'url': f"https://docs.google.com/spreadsheets/d/{sheet3_id}/edit"
            },
            'audit_logs': {
                'id': sheet4_id,
                'name': '4_CBEO_Audit_Backup_Logs',
                'url': f"https://docs.google.com/spreadsheets/d/{sheet4_id}/edit"
            }
        }
    }
    with open('drive_config.json', 'w', encoding='utf-8') as f:
        json.dump(config, f, indent=2, ensure_ascii=False)

    master_data = {
        'cbeo_info': {
            'office_name': 'कार्यालय मुख्य ब्लॉक शिक्षा अधिकारी (CBEO), भिनाय',
            'block': 'भिनाय (BHINAI)',
            'district': 'अजमेर / केकड़ी (AJMER)',
            'nic_sd_id': '8140',
            'ifms_id': '1408',
            'cbeo_officer': 'प्रमिला रासलोत',
            'mobile': '9414000000',
            'email': 'cbeo.bhinai.ajmer@rajasthan.gov.in'
        },
        'peeos': peeo_accounts,
        'staff': all_staff,
        'config': config
    }
    with open('master_cbeo_data.json', 'w', encoding='utf-8') as f:
        json.dump(master_data, f, indent=2, ensure_ascii=False)

    # Also save as master_cbeo_data.js for instant static / GitHub Pages zero-setup loading
    with open('master_cbeo_data.js', 'w', encoding='utf-8') as f:
        f.write("const MASTER_CBEO_DATA = " + json.dumps(master_data, indent=2, ensure_ascii=False) + ";\n")

    print("\nSUCCESS! Google Drive Backend Created and Master Data Exported.")
    print("Drive Folder URL:", config['folder_url'])
    for k, v in config['sheets'].items():
        print(f" - {v['name']}: {v['url']}")

if __name__ == '__main__':
    main()
