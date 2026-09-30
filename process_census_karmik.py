import openpyxl
import os
import sys
import json
import re

sys.path.append(r'C:\Users\jiten\Desktop\class11\hlb_detail\scratch')
from krutidev2unicode import kru2uni

CENSUS_FILE = r'C:\Users\jiten\Desktop\census2026\New folder\Bhinai Main Census 2026 Revised after clean.xlsx'
PRINCIPAL_FILE = 'school principal bhinai.xlsx'

def clean_key(s):
    if not s:
        return ''
    return re.sub(r'[^a-z0-9]', '', str(s).lower())

def convert_to_unicode(text):
    if not text:
        return ''
    s = str(text).strip()
    if not s or s.lower() == 'none':
        return ''
    
    # If already pure devanagari without KrutiDev letters
    # KrutiDev uses ASCII letters [a-zA-Z] to represent devanagari glyphs
    has_ascii_letters = bool(re.search(r'[a-zA-Z]', s))
    if not has_ascii_letters:
        return s.replace('¼', '(').replace('½', ')')

    # Don't convert if it's an SSO ID, code, or email
    if s.startswith(('RJ', 'STF', 'SBIN', 'BARB', 'IFSC', 'http')) or '@' in s:
        return s

    try:
        converted = kru2uni(s)
        converted = converted.replace('¼', '(').replace('½', ')').strip()
        # Clean common KrutiDev conversion artifacts
        converted = converted.replace('  ', ' ')
        return converted
    except Exception:
        return s

def main():
    print(f"Loading {CENSUS_FILE}...")
    wb = openpyxl.load_workbook(CENSUS_FILE, data_only=True)

    # 1. Build SSO Unicode Lookup from Superviser and Praganak
    sso_lookup = {}
    ws_sup = wb['Superviser']
    for r in range(1, ws_sup.max_row + 1):
        name = ws_sup.cell(r, 2).value
        post = ws_sup.cell(r, 5).value
        sso = ws_sup.cell(r, 7).value
        if sso and name:
            sso_lookup[str(sso).strip()] = {
                'name': str(name).strip(),
                'post': str(post).strip() if post else ''
            }

    ws_prag = wb['Praganak']
    for r in range(2, ws_prag.max_row + 1):
        name = ws_prag.cell(r, 4).value
        post = ws_prag.cell(r, 7).value
        sso = ws_prag.cell(r, 9).value
        if sso and name:
            sso_lookup[str(sso).strip()] = {
                'name': str(name).strip(),
                'post': str(post).strip() if post else ''
            }

    print(f"Preloaded {len(sso_lookup)} verified Unicode staff names from Superviser & Praganak sheets.")

    # 2. Load Master Data for 25 PEEOs
    with open('master_cbeo_data.json', 'r', encoding='utf-8') as f:
        master = json.load(f)

    peeos = master['peeos']

    # Pre-build PEEO lookup index
    # We map school names and village/panchayat names to PEEOs
    school_to_peeo = {}
    for p in peeos:
        p_name = p['peeo_name']
        school_to_peeo[clean_key(p['peeo_name'])] = p_name
        school_to_peeo[clean_key(p['panchayat_name'])] = p_name
        # extract core name e.g. "BANDANWARA", "DEOLIYA", etc.
        core = clean_key(p['peeo_name'].replace('PEEO', ''))
        if core:
            school_to_peeo[core] = p_name

        if p.get('schools'):
            for sch in p['schools']:
                school_to_peeo[clean_key(sch['school_name'])] = p_name
                school_to_peeo[clean_key(sch['village'])] = p_name

    # 3. Extract All Karmiks from Sheet1
    ws1 = wb['Sheet1']
    karmiks = []
    seen_sso = set()

    for r in range(4, ws1.max_row + 1):
        name_raw = ws1.cell(r, 2).value
        if not name_raw:
            continue

        sso = str(ws1.cell(r, 7).value or '').strip()
        if sso == 'None':
            sso = ''

        # Check lookup first for high-fidelity Unicode name
        if sso and sso in sso_lookup:
            name = sso_lookup[sso]['name']
            post = sso_lookup[sso]['post'] or convert_to_unicode(ws1.cell(r, 5).value)
        else:
            name = convert_to_unicode(name_raw)
            post = convert_to_unicode(ws1.cell(r, 5).value)

        gender = str(ws1.cell(r, 3).value or '').strip()
        dob = str(ws1.cell(r, 4).value or '').strip()
        school = convert_to_unicode(ws1.cell(r, 6).value)
        mob = str(ws1.cell(r, 8).value or '').strip()
        email = str(ws1.cell(r, 9).value or '').strip()
        bank_name = str(ws1.cell(r, 10).value or '').strip()
        bank_acc = str(ws1.cell(r, 11).value or '').strip()
        ifsc = str(ws1.cell(r, 12).value or '').strip()
        pan = str(ws1.cell(r, 13).value or '').strip()
        aadhaar = str(ws1.cell(r, 14).value or '').strip()

        # Clean mobile
        if mob.endswith('.0'):
            mob = mob[:-2]
        if mob in ['0', 'None', '-']:
            mob = ''

        # Clean email
        if email.lower() in ['none', '-']:
            email = ''

        # Map to PEEO
        ck_school = clean_key(school)
        matched_peeo = 'PEEO BHINAY' # default fallback
        
        # Exact/Partial matching
        for key, p_assigned in school_to_peeo.items():
            if len(key) >= 4 and (key in ck_school or ck_school in key):
                matched_peeo = p_assigned
                break

        staff_id = f"STF{1000 + len(karmiks) + 1}"

        karmiks.append({
            'staff_id': staff_id,
            'name': name,
            'gender': gender,
            'dob': dob,
            'post': post if post else 'अध्यापक / शिक्षक',
            'school_name': school if school else matched_peeo,
            'peeo_name': matched_peeo,
            'sso_id': sso,
            'mobile': mob,
            'email': email,
            'bank_name': bank_name if bank_name != 'None' else '',
            'bank_acc': bank_acc if bank_acc != 'None' else '',
            'ifsc': ifsc if ifsc != 'None' else '',
            'pan': pan if pan != 'None' else '',
            'aadhaar': aadhaar if aadhaar != 'None' else '',
            'status': 'Active'
        })

    print(f"Total Karmiks loaded from Census Master: {len(karmiks)}")

    # 4. Update master_data JSON & JS
    master['staff'] = karmiks
    master['total_staff_count'] = len(karmiks)

    with open('master_cbeo_data.json', 'w', encoding='utf-8') as f:
        json.dump(master, f, indent=2, ensure_ascii=False)

    with open('master_cbeo_data.js', 'w', encoding='utf-8') as f:
        f.write("const MASTER_CBEO_DATA = " + json.dumps(master, indent=2, ensure_ascii=False) + ";\n")

    print("Updated master_cbeo_data.json and master_cbeo_data.js successfully!")

    # 5. Sync to Google Drive Sheet: 2_CBEO_Staff_Directory
    print("\nUpdating Google Drive 2_CBEO_Staff_Directory sheet...")
    try:
        from create_cbeo_drive_backend import get_services
        drive, sheets = get_services()
        with open('drive_config.json', 'r', encoding='utf-8') as f:
            cfg = json.load(f)
        sheet2_id = cfg['sheets']['staff_directory']['id']

        s2_rows = [
            ['क्र.सं.', 'Staff ID', 'कार्मिक का नाम', 'पद', 'विद्यालय का नाम', 'संबंधित PEEO', 'मोबाइल नं.', 'ईमेल ID', 'SSO ID', 'बैंक खाता संख्या', 'IFSC कोड', 'PAN नं.', 'स्थिति']
        ]
        for idx, s in enumerate(karmiks, 1):
            s2_rows.append([
                idx,
                s['staff_id'],
                s['name'],
                s['post'],
                s['school_name'],
                s['peeo_name'],
                s['mobile'],
                s['email'],
                s['sso_id'],
                s['bank_acc'],
                s['ifsc'],
                s['pan'],
                s['status']
            ])

        # Clear and update sheet
        sheets.spreadsheets().values().clear(spreadsheetId=sheet2_id, range="Sheet1!A1:Z2000").execute()
        sheets.spreadsheets().values().update(
            spreadsheetId=sheet2_id,
            range="Sheet1!A1",
            valueInputOption="RAW",
            body={"values": s2_rows}
        ).execute()
        print(f"SUCCESS: Google Drive Sheet updated with {len(s2_rows)-1} karmiks!")
    except Exception as e:
        print("Note: Google Drive sync error (or offline):", e)

if __name__ == '__main__':
    main()
