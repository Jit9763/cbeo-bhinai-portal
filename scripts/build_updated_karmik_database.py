import openpyxl
import json
import re
import os

excel_path = r'C:\Users\jiten\Desktop\during census\14 june\Teacher Record_8140_07102026131201_6a542781-f.xlsx'
wb = openpyxl.load_workbook(excel_path)
s = wb['Sheet1']

with open('master_cbeo_data.json', encoding='utf-8') as f:
    master = json.load(f)

existing_staff = master.get('staff', [])
schools_56 = master.get('schools_56', [])

# 1. Existing verified Hindi names lookup
by_sso_hi = {}
by_name_en_hi = {}

for st in existing_staff:
    sso = (st.get('sso_id') or '').strip().upper()
    if sso and st.get('name') and not re.search(r'[a-zA-Z]', st.get('name')):
        by_sso_hi[sso] = st.get('name').strip()
    nen = (st.get('name_en') or '').strip().upper()
    if nen and st.get('name') and not re.search(r'[a-zA-Z]', st.get('name')):
        by_name_en_hi[nen] = st.get('name').strip()

# Verified manual Hindi names
MANUAL_NAME_HINDI = {
    'ABDUL AAHAD': 'अब्दुल आहाद',
    'ABDUL MOHSIN KHAN': 'अब्दुल मोहसिन खान',
    'ABHILASHA SUKHWAL': 'अभिलाषा सुखवाल',
    'ABHISHEK SHARMA': 'अभिषेक शर्मा',
    'AFSAR AHMED': 'अफसर अहमद',
    'AJAY SHARMA': 'अजय शर्मा',
    'AMIN MOHAMMAD': 'अमीन मोहम्मद',
    'AMITABH SANADHYA': 'अमिताभ सनाढ्य',
    'ANIL KUMAR SANKHLA': 'अनिल कुमार सांखला',
    'ANIL KUMAR SHARMA': 'अनिल कुमार शर्मा',
    'ANITA CHOUDHARY': 'अनीता चौधरी',
    'ANITA KAHAR': 'अनीता कहार',
    'ANITA SWAMI': 'अनीता स्वामी',
    'ANUPAMA DOSAYA': 'अनुपमा दोसाया',
    'ANUPAMA SANKHLA': 'अनुपमा सांखला',
    'ANURAG SINGHAL': 'अनुराग सिंघल',
    'ARCHANA JADOUN': 'अर्चना जादौन',
    'ASHISH ASWAL': 'आशीष असवाल',
    'ASHISH CHOUDHARY': 'आशीष चौधरी',
    'ASHISH VERMA': 'आशीष वर्मा',
    'ASHOK KUMAR JAT': 'अशोक कुमार जाट',
    'ASHWINI KUMAR': 'अश्विनी कुमार',
    'BABITA KANWAR SHEKHAWAT': 'बबीता कंवर शेखावत',
    'BADRI PRASAD DEORA': 'बद्री प्रसाद देवड़ा',
    'BAL RAM YADAV': 'बलराम यादव',
    'BALBEER SINGH YADAV': 'बलबीर सिंह यादव',
    'BALMUKAND SHARMA': 'बालमुकुंद शर्मा',
    'BALVEER UJJAWAL': 'बलवीर उज्ज्वल',
    'BANWARI MALI': 'बनवारी माली',
    'BEENA KUMARI MEENA': 'बीना कुमारी मीणा',
    'BHAGCHAND JAIN': 'भागचंद जैन',
    'BHAGCHAND REGAR': 'भागचंद रैगर',
    'BHAGCHAND  TANK': 'भागचंद टांक',
    'BHAIRON SINGH': 'भैरों सिंह',
    'BHANWAR LAL JAT': 'भंवर लाल जाट',
    'BHERU LAL BALAI': 'भेरू लाल बलाई',
    'BIJENDRA KUMAR': 'बिजेंद्र कुमार',
    'CHANDAR SHAKAR SHARMA': 'चंद्रशेखर शर्मा',
    'CHENA RAM': 'चेनाराम',
    'CHETAN KUMAR CHAUDHARY': 'चेतन कुमार चौधरी',
    'DEBI SINGH SHEKHAWAT': 'देवी सिंह शेखावत',
    'DEEPAK KUMAR JOSHI': 'दीपक कुमार जोशी',
    'DEEPAK SANWARIYA': 'दीपक सांवरिया',
    'DEEPCHAND MEENA': 'दीपचंद मीणा',
    'DEVENDRA KUMAR MISHRA': 'देवेन्द्र कुमार मिश्रा',
    'DEVENDRA PRAJAPAT': 'देवेन्द्र प्रजापत',
    'DEVI LAL CHANDEL': 'देवी लाल चंदेल',
    'DEVRAJ GURJAR': 'देवराज गुर्जर',
    'DHANNE SINGH RAO': 'धन्ने सिंह राव',
    'DHANRAJ SAINI': 'धनराज सैनी',
    'DHARMENDRA JHAROTIYA': 'धर्मेन्द्र झारोटिया',
    'DHARMRAJ MEENA': 'धर्मराज मीणा',
    'DINESH KUMAR': 'दिनेश कुमार',
    'DINESH KUMAR KUMAWAT': 'दिनेश कुमार कुमावत',
    'GAGANDEEP SINGH RATHORE': 'गगनदीप सिंह राठौड़',
    'GAJRAJ SINGH': 'गजराज सिंह',
    'GIRDHAR GOPAL BALAI': 'गिरधर गोपाल बलाई',
    'GIRIRAJ CHOUDHARY': 'गिरिराज चौधरी',
    'GOVIND LAL': 'गोविंद लाल',
    'GULSHAN BHARTI': 'गुलशन भारती',
    'GULSHAN SONWAR': 'गुलशन सोनवार',
    'HANS RAJ GURJAR': 'हंसराज गुर्जर',
    'HARENDRA KUMAR': 'हरेन्द्र कुमार',
    'HARI KISHAN BUNKER': 'हरीकिशन बुनकर',
    'HARICHHA BALOTIA': 'हरीच्छा बालोटिया',
    'HARISH KUMAR VAISHNAV': 'हरीश कुमार वैष्णव',
    'HARISH SAGAR': 'हरीश सागर',
    'HEMLATA SANKHLA': 'हेमलता सांखला',
    'HEMRAJ DAROGA': 'हेमराज दरोगा',
    'HEMRAJ SAINI': 'हेमराज सैनी',
    'HINA BANO': 'हिना बानो',
    'HUSAIN MOHAMMAD SIPAI': 'हुसैन मोहम्मद सिपाही',
    'INDRA KUMARI BALAI': 'इन्द्रा कुमारी बलाई',
    'JAGDISH CHOUDHARY': 'जगदीश चौधरी',
    'JAI PRAKASH JANGIR': 'जय प्रकाश जांगिड़',
    'JALAJ TAK': 'जलज टाक',
    'JAYA  PARIHAR': 'जया परिहार',
    'JAYA PARIHAR': 'जया परिहार',
    'JAYA TOLANI': 'जया तोलानी',
    'JEEVANRAM BAIRWA': 'जीवनराम बैरवा',
    'KALILA BEGUM': 'कलीला बेगम',
    'KAMLESH CHAND': 'कमलेश चंद',
    'KAVITA KUMARI MEENA': 'कविता कुमारी मीणा',
    'KHUSHBOO SHARMA': 'खुशबू शर्मा',
    'KISHOR KUMAR SWAMI': 'किशोर कुमार स्वामी',
    'KOMAL SINGH MEENA': 'कोमल सिंह मीणा',
    'KULDEEP MISHRA': 'कुलदीप मिश्रा',
    'KULDEEP SINGH': 'कुलदीप सिंह',
    'KUTUBUDDEEN': 'कुतुबुद्दीन',
    'LAKHAN VAISHNAV': 'लखन वैष्णव',
    'MADHYAMA VASHISHTH': 'मध्यमा वशिष्ठ',
    'MAHAVEER SINGH DEVRA': 'महावीर सिंह देवड़ा',
    'MAHIMA JOSHI': 'महिमा जोशी',
    'MAHIPAL SINGH CHOUHAN': 'महिपाल सिंह चौहान',
    'MANISH ACHARYA': 'मनीष आचार्य',
    'MANISHA JANGID': 'मनीषा जांगिड़',
    'MANJU GURJAR': 'मंजू गुर्जर',
    'MANOJ KUMAR MEENA': 'मनोज कुमार मीणा',
    'MARUDHAR RATHORE': 'मरुधर राठौड़',
    'MEENU SAINI': 'मीनू सैनी',
    'MINAKSHI SAINI': 'मीनाक्षी सैनी',
    'MINTU KANWAR': 'मिंटू कंवर',
    'MOHAMMAD JAVED GOURI': 'मोहम्मद जावेद गौरी',
    'MOHAMMED SHAREEF MOMINE': 'मोहम्मद शरीफ मोमिन',
    'MOHAN LAL SAIN': 'मोहन लाल सैन',
    'MONA JANGID': 'मोना जांगिड़',
    'MONIKA BAI MEENA': 'मोनिका बाई मीणा',
    'MONIKA JANGID': 'मोनिका जांगिड़',
    'MUDITA YADAV': 'मुदिता यादव',
    'MUKESH MEENA': 'मुकेश मीणा',
    'NAVEEN SINGH RAWAT': 'नवीन सिंह रावत',
    'NAVEL KISHORE': 'नवल किशोर',
    'NAVNEET KUMAR MISHRA': 'नवनीत कुमार मिश्रा',
    'NAZMEEN': 'नजमीन',
    'NEELAM PAREEK': 'नीलम पारीक',
    'NISHA MEENA': 'निशा मीणा',
    'NITA SHUKLA': 'नीता शुक्ला',
    'OM PRAKASH CHOUDHRY': 'ओम प्रकाश चौधरी',
    'OMPRAKASH SERAWAT': 'ओमप्रकाश सेरावत',
    'PARSA RAM': 'परसाराम',
    'PAWAN KUMAR GURJAR': 'पवन कुमार गुर्जर',
    'PINKY SUMMERWAR': 'पिंकी समरवार',
    'PRAHLAD DASS BHAMBHU': 'प्रहलाद दास भांभू',
    'PRAKASH TIWARI': 'प्रकाश तिवारी',
    'PRAMOD KUMAR SHARMA': 'प्रमोद कुमार शर्मा',
    'PRATIBHA TAILOR': 'प्रतिभा टेलर',
    'PRAVEEN KUMAR SHARMA': 'प्रवीण कुमार शर्मा',
    'PREMLATA CHOUHAN': 'प्रेमलता चौहान',
    'PUNAM CHAND TAK': 'पूनम चंद टाक',
    'PURNIMA': 'पूर्णिमा',
    'PURNIMA KUMAWAT': 'पूर्णिमा कुमावत',
    'PUSHPENDRA CHAUHAN': 'पुष्पेंद्र चौहान',
    'RAGHAV MAHESWARI': 'राघव माहेश्वरी',
    'RAJEEV KUMAR MANDOT': 'राजीव कुमार मांडोत',
    'RAJENDAR KUMAR': 'राजेंद्र कुमार',
    'RAJENDRA KHATI': 'राजेंद्र खाती',
    'RAJESH KUMAR REWALA': 'राजेश कुमार रेवाला',
    'RAJU KUMARI MEENA': 'राजू कुमारी मीणा',
    'RAJU LAL NAYAK': 'राजू लाल नायक',
    'RAJU LAL REBARI': 'राजू लाल रेबारी',
    'RAKESH OJHA': 'राकेश ओझा',
    'RAM PRASAD REGAR': 'रामप्रसाद रैगर',
    'RAMCHANDRA JAT': 'रामचंद्र जाट',
    'RAMDEV BAIRWA': 'रामदेव बैरवा',
    'RAMKUMARI MEENA': 'रामकुमारी मीणा',
    'RAMSWAROOP BAIRWA': 'रामस्वरूप बैरवा',
    'RANG LAL GURJAR': 'रंगलाल गुर्जर',
    'RINKU JANGID': 'रिंकू जांगिड़',
    'SAMEEM BANOO': 'शमीम बानो',
    'SANDEEP KASOTIYA': 'संदीप कसोतिया',
    'SANJAY KUMAR': 'संजय कुमार',
    'SANTOSH DEVI YADAV': 'संतोष देवी यादव',
    'SARITA MOOND': 'सरिता मूंड',
    'SHAHANA': 'शहाना',
    'SHAHRUKH KHAN': 'शाहरुख खान',
    'SHAMBHU NATH JOGI': 'शंभू नाथ योगी',
    'SHANKAR LAL JAT': 'शंकर लाल जाट',
    'SHARIF MOHAMMED BHATI': 'शरीफ मोहम्मद भाटी',
    'SHEELA CHOUHAN': 'शीला चौहान',
    'SHIV JI GURJAR': 'शिवजी गुर्जर',
    'SHIVANSH KUMAR BAIRAGI': 'शिवांश कुमार बैरागी',
    'SHIVRAJ PRAJAPAT': 'शिवराज प्रजापत',
    'SHUBHAM CHOUHAN': 'शुभम चौहान',
    'SHYAM SINGH SHEKHAWAT': 'श्याम सिंह शेखावत',
    'SITA CHOUDHARY': 'सीता चौधरी',
    'SOHAN LAL SARAN': 'सोहनलाल सारण',
    'SOMESHWAR BAKOLIYA': 'सोमेश्वर बाकोलिया',
    'SONAM SAINI': 'सोनम सैनी',
    'SONU KUMARI CHOUDHARY': 'सोनू कुमारी चौधरी',
    'SUBE SINGH JAT': 'सूबे सिंह जाट',
    'SUBHASH CHAND BIJARNIYA': 'सुभाष चंद बिजारणिया',
    'SUDARSHAN KUMAR': 'सुदर्शन कुमार',
    'SUKHA RAM': 'सुखाराम',
    'SUNIL KUMAR MEGHWANSHI': 'सुनील कुमार मेघवंशी',
    'SUNIL KUMAR SASHI': 'सुनील कुमार शशि',
    'SUNITA MAHAN': 'सुनीता महान',
    'SUNITA MARU': 'सुनीता मारू',
    'SURAJ GURJAR': 'सूरज गुर्जर',
    'SURENDAR CHHAWRI': 'सुरेंद्र छावड़ी',
    'SURENDRA KUMAR MANGAVA': 'सुरेंद्र कुमार मंगवा',
    'SUSHMA YADAV': 'सुषमा यादव',
    'TARUNA SOLANKI': 'तरुणा सोलंकी',
    'USHA FADOLIYA': 'उषा फड़ोलिया',
    'VAIBHAV RAJ MEHRA': 'वैभव राज मेहरा',
    'VARSHA YADAV': 'वर्षा यादव',
    'VIJENDRA MEENA': 'विजेंद्र मीणा',
    'VIKARANT VAISHNAV': 'विक्रांत वैष्णव',
    'VINOD BHAMBI': 'विनोद भांभी',
    'VINOD KUMAR': 'विनोद कुमार',
    'VINOD KUMAR MEENA': 'विनोद कुमार मीणा',
    'VIPIN KUMAR': 'विपिन कुमार',
    'VISHAL YADUVANSHI': 'विशाल यदुवंशी',
    'VISHNU GURJAR': 'विष्णु गुर्जर',
    'VISHNU JANGID': 'विष्णु जांगिड़',
    'VIVEK CHOUDHARY': 'विवेक चौधरी',
    'YADUVEER SINGH': 'यदुवीर सिंह',
    'YOGENDRA PRATAP SINGH': 'योगेंद्र प्रताप सिंह'
}

# Designation Map
POST_MAP = {
    'Principal and Equivalent': ('प्रधानाचार्य', 'Principal'),
    'Vice Principal (School)': ('उप प्रधानाचार्य', 'Vice Principal'),
    'Headmaster and Equivalent': ('प्रधानाध्यापक', 'Headmaster'),
    'Lecturer (I Gr.)': ('प्राध्यापक (स्कूल शिक्षा)', 'Lecturer (I Gr.)'),
    'Lecturer of Special Education': ('व्याख्याता (विशेष शिक्षा)', 'Lecturer of Special Education'),
    'Senior Teacher': ('वरिष्ठ अध्यापक', 'Senior Teacher'),
    'Teacher Level-2': ('अध्यापक लेवल-2', 'Teacher Level-2'),
    'Teacher Level-2 (Special Education)': ('अध्यापक लेवल-2 (विशेष शिक्षा)', 'Teacher Level-2 (Spl. Ed.)'),
    'Teacher Level-1': ('अध्यापक लेवल-1', 'Teacher Level-1'),
    'Teacher Level-1 (Special Education)': ('अध्यापक लेवल-1 (विशेष शिक्षा)', 'Teacher Level-1 (Spl. Ed.)'),
    'Teacher (III Gr.) Level 1 / Teacher (III Gr.)': ('अध्यापक लेवल-1', 'Teacher Level-1'),
    'Assistant Teacher(L-2)': ('सहायक अध्यापक लेवल-2', 'Assistant Teacher (L-2)'),
    'Pre Primary Teacher (L-8)': ('पूर्व प्राथमिक शिक्षक (L-8)', 'Pre Primary Teacher (L-8)'),
    'Basic Computer Instructor': ('बेसिक कंप्यूटर अनुदेशक', 'Basic Computer Instructor'),
    'Physical Education Teacher': ('शारीरिक शिक्षक', 'Physical Education Teacher (PET)'),
    'Senior Physical Education Teacher': ('वरिष्ठ शारीरिक शिक्षक', 'Senior PET'),
    'Prabodhak Level-1': ('प्रबोधक लेवल-1', 'Prabodhak Level-1'),
    'Prabodhak Level-2': ('प्रबोधक लेवल-2', 'Prabodhak Level-2'),
    'Prabodhak Sharirik Shikshak': ('प्रबोधक शारीरिक शिक्षक', 'Prabodhak PET'),
    'Varishth prabodhak': ('वरिष्ठ प्रबोधक', 'Senior Prabodhak'),
    'Varishth prabodhak -': ('वरिष्ठ प्रबोधक', 'Senior Prabodhak'),
    'Librarian (II Gr.)': ('पुस्तकालयाध्यक्ष द्वितीय श्रेणी', 'Librarian (II Gr.)'),
    'Librarian (III Gr.)': ('पुस्तकालयाध्यक्ष तृतीय श्रेणी', 'Librarian (III Gr.)'),
    'Lab Assistant': ('प्रयोगशाला सहायक', 'Lab Assistant'),
    'Additional Administrative Officer': ('अतिरिक्त प्रशासनिक अधिकारी', 'Additional Admin Officer'),
    'Assistant Administrative Officer': ('सहायक प्रशासनिक अधिकारी', 'Assistant Admin Officer'),
    'Senior Assistant': ('वरिष्ठ सहायक', 'Senior Assistant'),
    'Junior Assistant': ('कनिष्ठ सहायक', 'Junior Assistant'),
    'Class IV': ('चतुर्थ श्रेणी कर्मचारी', 'Class IV Staff'),
    '': ('अध्यापक', 'Teacher')
}

# 2. School and PEEO lookups
school_to_peeo = {}
school_code_to_hi = {}
school_code_to_en = {}

for p in master['peeos']:
    p_name = p['peeo_name']
    for sch in p.get('schools', []):
        code = str(sch.get('shala_darpan_code') or sch.get('dise_code') or '').strip()
        if code:
            school_to_peeo[code] = p_name
            if sch.get('school_name'):
                school_code_to_hi[code] = sch.get('school_name')

for sch in schools_56:
    code = str(sch.get('shala_darpan_code') or '').strip()
    if code:
        school_to_peeo[code] = sch.get('peeo_name')
        if sch.get('school_name_hi'):
            school_code_to_hi[code] = sch.get('school_name_hi')
        if sch.get('school_name_en'):
            school_code_to_en[code] = sch.get('school_name_en')

# Specific fallback for sanskrit schools
school_to_peeo['221779'] = 'PEEO BHINAY'
school_to_peeo['226359'] = 'PEEO DEOLIYA KALAN'
school_to_peeo['226361'] = 'PEEO DHANTOL'
school_to_peeo['226360'] = 'PEEO NAGOLA'

# Verified Principals lookup from schools_56
verified_prins_map = {}
for sch in schools_56:
    code = str(sch.get('shala_darpan_code') or '').strip()
    if code and sch.get('principal_name'):
        verified_prins_map[code] = {
            'name': sch.get('principal_name', '').strip(),
            'mobile': sch.get('principal_mobile', '').strip(),
            'school_name_hi': sch.get('school_name_hi', '')
        }

# Process all rows in new Excel
new_staff_list = []
seen_school_heads = set()

for r in range(2, s.max_row+1):
    sch_code = str(s.cell(row=r, column=1).value or '').strip()
    sch_raw = str(s.cell(row=r, column=7).value or '').strip()
    m_code = re.search(r'\((\d+)\)', sch_raw)
    clean_sch_name_en = sch_raw[:m_code.start()].strip() if m_code else sch_raw
    
    post_raw = str(s.cell(row=r, column=19).value or '').strip()
    subject_raw = str(s.cell(row=r, column=20).value or '').strip()
    name_raw = str(s.cell(row=r, column=21).value or '').strip()
    emp_id = str(s.cell(row=r, column=22).value or '').strip().upper()
    nic_id = str(s.cell(row=r, column=23).value or '').strip()
    gender_raw = str(s.cell(row=r, column=24).value or '').strip()
    dob_raw = str(s.cell(row=r, column=30).value or '').strip()
    joining_raw = str(s.cell(row=r, column=31).value or '').strip()
    mob_raw = str(s.cell(row=r, column=48).value or '').strip()
    email_raw = str(s.cell(row=r, column=49).value or '').strip()
    retire_raw = str(s.cell(row=r, column=60).value or '').strip()
    status_raw = str(s.cell(row=r, column=68).value or 'Working').strip()
    
    # 1. Hindi Name resolution
    clean_en_upper = re.sub(r'\s+', ' ', name_raw).strip().upper()
    name_hi = by_sso_hi.get(emp_id) or by_name_en_hi.get(clean_en_upper) or MANUAL_NAME_HINDI.get(clean_en_upper)
    if not name_hi:
        name_hi = clean_en_upper.title()  # fallback
    
    # English Name in clean Title Case
    name_en = clean_en_upper.title()

    # 2. Post / Designation resolution
    post_hi, post_en = POST_MAP.get(post_raw, (post_raw, post_raw))

    # 3. School and PEEO
    peeo_name = school_to_peeo.get(sch_code) or f"PEEO {str(s.cell(row=r, column=4).value or 'BHINAY').strip().upper()}"
    sch_name_hi = school_code_to_hi.get(sch_code) or clean_sch_name_en
    sch_name_en = school_code_to_en.get(sch_code) or clean_sch_name_en

    # 4. Check if Principal / Head of School
    is_pradhan = False
    is_principal_post = (post_raw == 'Principal and Equivalent' or 'Principal' in post_raw)
    
    ver_p = verified_prins_map.get(sch_code)
    if ver_p:
        v_name_clean = ver_p['name'].lower()
        # check name similarity
        if any(part in clean_en_upper.lower() for part in v_name_clean.split() if len(part) > 2):
            is_pradhan = True
            if not mob_raw and ver_p['mobile']:
                mob_raw = ver_p['mobile']
        elif is_principal_post and sch_code not in seen_school_heads:
            is_pradhan = True
    elif is_principal_post:
        is_pradhan = True
    
    if is_pradhan:
        seen_school_heads.add(sch_code)

    staff_id = emp_id if (emp_id and emp_id != 'NONE') else (f"NIC_{nic_id}" if nic_id else f"STF_{sch_code}_{r}")

    staff_obj = {
        "staff_id": staff_id,
        "name": name_hi,
        "name_en": name_en,
        "gender": "महिला" if gender_raw.lower() == "female" else "पुरुष",
        "dob": dob_raw,
        "post": post_hi,
        "post_en": post_en,
        "subject": subject_raw if subject_raw != '---' else '',
        "subject_en": subject_raw if subject_raw != '---' else '',
        "school_name": sch_name_hi,
        "school_name_hi": sch_name_hi,
        "school_name_en": sch_name_en,
        "shala_darpan_code": sch_code,
        "school_code": sch_code,
        "peeo_name": peeo_name,
        "sso_id": emp_id,
        "employee_id": emp_id,
        "nic_id": nic_id,
        "mobile": mob_raw,
        "email": email_raw if email_raw != 'None' else '',
        "joining_date": joining_raw,
        "retirement_date": retire_raw,
        "is_sanstha_pradhan": is_pradhan,
        "status": "Active" if status_raw.lower() in ['working', 'permanent'] else "Active",
        "remarks": "सत्यापित संस्था प्रधान" if is_pradhan else "",
        "unmapped_school": False
    }
    new_staff_list.append(staff_obj)

print(f"Total processed new staff records: {len(new_staff_list)}")
print(f"Identified Sanstha Pradhans in list: {sum(1 for s in new_staff_list if s['is_sanstha_pradhan'])}")

# Update master_cbeo_data.json
master['staff'] = new_staff_list
master['total_staff_count'] = len(new_staff_list)
if 'cbeo_info' in master:
    master['cbeo_info']['total_staff'] = len(new_staff_list)

with open('master_cbeo_data.json', 'w', encoding='utf-8') as f:
    json.dump(master, f, ensure_ascii=False, indent=2)
print("Updated master_cbeo_data.json successfully!")

# Update master_cbeo_data.js
js_content = 'const MASTER_CBEO_DATA = ' + json.dumps(master, ensure_ascii=False, indent=2) + ';\n'
with open('master_cbeo_data.js', 'w', encoding='utf-8') as f:
    f.write(js_content)
print("Updated master_cbeo_data.js successfully!")
