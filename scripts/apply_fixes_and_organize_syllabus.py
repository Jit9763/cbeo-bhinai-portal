# -*- coding: utf-8 -*-
import os
import shutil
import openpyxl
import json
from PIL import Image

print("--- STEP 1: STANDARDIZE VILLAGE NAMES (CHHACHHUNDRA & GHANA) ---")

text_files = [
    'master_cbeo_data.json',
    'master_cbeo_data.js',
    'schools_56_master.json',
    'scripts/latest_report.md',
    'scripts/saman_enrolment_match_report.html',
    'scripts/standardize_villages_and_schools.py'
]

for fp in text_files:
    if os.path.exists(fp):
        with open(fp, 'r', encoding='utf-8') as f:
            content = f.read()
        new_content = content.replace('छाछून्दरा', 'छछून्दरा').replace('घाणा', 'घणा')
        if new_content != content:
            with open(fp, 'w', encoding='utf-8') as f:
                f.write(new_content)
            print(f"Updated: {fp}")

# Also update school enrolment.xlsx
wb_path = 'school enrolment.xlsx'
if os.path.exists(wb_path):
    wb = openpyxl.load_workbook(wb_path)
    updated_cells = 0
    for sname in wb.sheetnames:
        sheet = wb[sname]
        for r in range(1, sheet.max_row + 1):
            for c in range(1, sheet.max_column + 1):
                cell = sheet.cell(row=r, column=c)
                if isinstance(cell.value, str):
                    if 'छाछून्दरा' in cell.value or 'घाणा' in cell.value:
                        new_val = cell.value.replace('छाछून्दरा', 'छछून्दरा').replace('घाणा', 'घणा')
                        cell.value = new_val
                        updated_cells += 1
    wb.save(wb_path)
    print(f"Updated school enrolment.xlsx: {updated_cells} cells standardized.")

print("\n--- STEP 2: CREATE DEDICATED FOLDER FOR PENDING SYLLABUS FILES ---")

syllbus_dir = r"C:\Users\jiten\Desktop\cbeo\syllbus"
dest_folder = os.path.join(syllbus_dir, "लंबित_पोर्टल_पर_नहीं_आए")
os.makedirs(dest_folder, exist_ok=True)

pending_matches = [
    {
        "code": "221764",
        "name": "रा.उ.मा.वि. बड़गांव (सूरखण्ड)",
        "sources": [
            "WhatsApp Image 2026-10-07 at 8.13.58 AM.jpeg",
            "WhatsApp Image 2026-10-07 at 8.14.04 AM.jpeg"
        ],
        "dest_name": "221764_रा.उ.मा.वि._बड़गांव_सूरखण्ड_Syllabus"
    },
    {
        "code": "221770",
        "name": "महात्मा गांधी राजकीय विद्यालय, बांदनवाड़ा",
        "sources": [
            "पाठ्यक्रम पूर्णता प्रतिशत रिपोर्ट (Black & White).pdf"
        ],
        "dest_name": "221770_MGGS_बांदनवाड़ा_Syllabus"
    },
    {
        "code": "488791",
        "name": "रा.बा.उ.मा.वि. खेड़ी",
        "sources": [
            "GGSSS khedi पाठयक्रम पूर्णता सूचना.pdf",
            "WhatsApp Image 2026-10-07 at 8.13.59 AM.jpeg"
        ],
        "dest_name": "488791_रा.बा.उ.मा.वि._खेड़ी_Syllabus"
    },
    {
        "code": "488897",
        "name": "रा.उ.मा.वि. घणा",
        "sources": [
            "New Doc 10-06-2026 12.49.pdf"
        ],
        "dest_name": "488897_रा.उ.मा.वि._घणा_Syllabus"
    },
    {
        "code": "410632",
        "name": "रा.बा.उ.मा.वि. नांदसी",
        "sources": [
            "कोर्स GGSSS NANDSI.pdf"
        ],
        "dest_name": "410632_रा.बा.उ.मा.वि._नांदसी_Syllabus"
    },
    {
        "code": "221765",
        "name": "रा.उ.मा.वि. कनाई कलां",
        "sources": [
            "WhatsApp Image 2026-10-07 at 8.14.01 AM.jpeg"
        ],
        "dest_name": "221765_रा.उ.मा.वि._कनाई_कलां_Syllabus"
    },
    {
        "code": "488947",
        "name": "रा.उ.मा.वि. हियालिया",
        "sources": [
            "WhatsApp Image 2026-10-07 at 8.14.05 AM.jpeg"
        ],
        "dest_name": "488947_रा.उ.मा.वि._हियालिया_Syllabus"
    }
]

for item in pending_matches:
    code = item["code"]
    name = item["name"]
    base_dest = os.path.join(dest_folder, item["dest_name"])
    
    # If source is pdf, copy pdf
    # If source is jpeg, copy jpeg AND also convert to pdf for easy 1-click viewing
    img_list = []
    for sfile in item["sources"]:
        src_path = os.path.join(syllbus_dir, sfile)
        if not os.path.exists(src_path):
            print(f"Warning: {src_path} not found!")
            continue
        ext = os.path.splitext(sfile)[1].lower()
        if ext == '.pdf':
            dest_file = f"{base_dest}.pdf"
            shutil.copy2(src_path, dest_file)
            print(f"Copied PDF: {sfile} -> {os.path.basename(dest_file)}")
        elif ext in ['.jpeg', '.jpg', '.png']:
            dest_file = f"{base_dest}_{os.path.basename(sfile)}"
            shutil.copy2(src_path, dest_file)
            print(f"Copied Image: {sfile} -> {os.path.basename(dest_file)}")
            img_list.append(src_path)
            
    if img_list:
        # Generate combined/single PDF from image(s) for quick viewing!
        images = [Image.open(img).convert('RGB') for img in img_list]
        pdf_out = f"{base_dest}.pdf"
        images[0].save(pdf_out, save_all=True, append_images=images[1:])
        print(f"Generated clean PDF for user: {os.path.basename(pdf_out)}")

print(f"\nAll 7 pending schools organized into: {dest_folder}")
