# -*- coding: utf-8 -*-
import openpyxl
import json
import re

# 1. Load 97 standard villages from census
wb_census = openpyxl.load_workbook(r'C:\Users\jiten\Desktop\census2026\Village List Format Bhinai 1.xlsx')
s_census = wb_census['Sheet1']

census_97 = []
census_by_norm_en = {}
census_by_norm_hi = {}

for r in range(2, s_census.max_row + 1):
    sr = s_census.cell(row=r, column=1).value
    v_en = str(s_census.cell(row=r, column=16).value or '').strip()
    v_hi = str(s_census.cell(row=r, column=17).value or '').strip()
    lg = s_census.cell(row=r, column=18).value
    gp_hi = str(s_census.cell(row=r, column=20).value or '').strip()
    
    if v_hi:
        item = {'sr': sr, 'v_en': v_en, 'v_hi': v_hi, 'lg': lg, 'gp_hi': gp_hi}
        census_97.append(item)
        norm_en = re.sub(r'[^A-Z0-9]', '', v_en.upper())
        census_by_norm_en[norm_en] = item
        census_by_norm_hi[v_hi] = item

print(f"Loaded {len(census_97)} census villages.")

# 2. Comprehensive Alias Map for Village Strings found in School Master
VILLAGE_ALIASES = {
    'AMARGADH': 'अमरगढ',
    'AMARGARH': 'अमरगढ',
    'BADANVADA': 'बान्दनवाडा',
    'BANDANWARA': 'बान्दनवाडा',
    'BADGAON': 'बड़गांव',
    'BARGAON': 'बड़गांव',
    'BADLA KHEDA': 'बडला',
    'BADLA URF KALA TALAB': 'बडला',
    'BADLA': 'बडला',
    'BADLI': 'बड़ली',
    'BARLI': 'बड़ली',
    'BAGRAI (GURHA KHURD)': 'बगराई',
    'BAGRAI (KUMHARIYA)': 'बगराई',
    'BAGRAI': 'बगराई',
    'BAGARI': 'बगराई',
    'BALAPURA (EKALSEENGA)': 'बालापुरा',
    'BALAPURA (NAGOLA)': 'बालापुरा',
    'BALAPURA': 'बालापुरा',
    'BHINAY': 'भिनाय',
    'BHINAI': 'भिनाय',
    'BUVKIYA': 'बूबकिया',
    'BOOBKIYA': 'बूबकिया',
    'CHACHUNDRA': 'छछून्दरा',
    'CHHACHHUNDRA': 'छछून्दरा',
    'CHAMPANERI': 'चापानेरी',
    'CHAPANERI': 'चापानेरी',
    'DEOLIYA KALAN': 'देवलिया कलां',
    'DEVLIYA KALAN': 'देवलिया कलां',
    'DHANDHO KA KHERA': 'धांधो का खेड़ा',
    'DOLATPURA': 'दौलतपुरा',
    'EKALSEENGA': 'एकलसिंघा',
    'EKALSINGHA': 'एकलसिंघा',
    'GAJJA NADI': 'गज्जनाड़ी',
    'GAJJANAADI': 'गज्जनाड़ी',
    'GANAHERA': 'गनाहेड़ा',
    'GHANA': 'घणा',
    'GORDHANPURA': 'गोरधनपुरा',
    'GUDHA KHURD': 'गुढ़ा खुर्द',
    'GURHA KHURD': 'गुढ़ा खुर्द',
    'HARPURA': 'हरपुरा',
    'HEERAPURA': 'हीरापुरा',
    'HIYALIYA': 'हियालिया',
    'INDRAPURA': 'इन्द्रपुरा',
    'JAITPURA': 'जैतपुरा',
    'JETPURA': 'जैतपुरा',
    'JHIPIYA': 'झींपिया',
    'KADOLAI': 'कादोलाई',
    'KAIROT': 'कैरोट',
    'KEROT': 'कैरोट',
    'KANAI KALA': 'कनाई कलां',
    'KANAI KALAN': 'कनाई कलां',
    'KARANTI': 'करांटी',
    'KARATI': 'करांटी',
    'KHEDI': 'खेड़ी',
    'KUMHARIYA': 'कुम्हारिया',
    'KURTHAL': 'कूरथल',
    'LAMGARA': 'लाम्गरा',
    'NAGOLA': 'नागोला',
    'NANDSI': 'नांदसी',
    'NEMEDA': 'निमेड़ा',
    'NIMEDA': 'निमेड़ा',
    'PADANGA': 'पाडंगा',
    'PADLIYA': 'पाडलिया',
    'PADALIYA': 'पाडलिया',
    'PIPLIYA': 'पीपलिया',
    'RAGHUNATHGADH': 'रघुनाथगढ़',
    'RAMMALIYA': 'राममालिया',
    'RAMMALIA': 'राममालिया',
    'RATAKOT': 'राताकोट',
    'ROOPPURA': 'रूपपुरा',
    'SEDRIYA': 'सेदरिया',
    'SINGAWAL': 'सिंगावल',
    'SOBRI': 'सोबड़ी',
    'SOORAJPURA': 'सूरजपुरा',
    'SURKHAND': 'सूरखण्ड',
    'TANTOTI': 'टांटोटी',
    'KHERI': 'खेड़ी',
    'MOTIPURA': 'मोतीपुरा',
    'RAMPURA': 'रामपुरा',
    'KITAP': 'कीटाप',
    'BHAGWANPURA': 'भगवानपुरा'
}

# 3. Load Official Standard School Names from school enrolment.xlsx
wb_sd = openpyxl.load_workbook('school enrolment.xlsx')
s_sd = wb_sd['समान परीक्षा मिलान व अंतर']

std_school_names = {}
for r in range(7, 56):
    code = str(s_sd.cell(row=r, column=2).value or '').strip()
    raw_name = str(s_sd.cell(row=r, column=3).value or '').strip()
    if code and code.isdigit() and raw_name:
        # Split Hindi & English
        m = re.match(r'^(.*?)\s*\((.*?)\)$', raw_name)
        if m:
            hi_name = m.group(1).strip()
            en_name = m.group(2).strip()
        else:
            hi_name = raw_name
            en_name = raw_name
        std_school_names[code] = {
            'code': code,
            'name_hi': hi_name,
            'name_en': en_name,
            'full_official': raw_name
        }

print(f"Loaded {len(std_school_names)} standard school names from school enrolment.xlsx.")

# 4. Update master_cbeo_data.json
with open('master_cbeo_data.json', encoding='utf-8') as f:
    master = json.load(f)

# Update schools_56
for s in master.get('schools_56', []):
    c = str(s.get('shala_darpan_code', '')).strip()
    if c in std_school_names:
        s['school_name_hi'] = std_school_names[c]['name_hi']
        s['school_name_en'] = std_school_names[c]['name_en']
        s['school_name'] = std_school_names[c]['name_hi']

# Update all schools in peeos
total_updated_villages = 0
for p in master.get('peeos', []):
    for sch in p.get('schools', []):
        c = str(sch.get('shala_darpan_code', '')).strip()
        if c in std_school_names:
            sch['school_name_hi'] = std_school_names[c]['name_hi']
            sch['school_name_en'] = std_school_names[c]['name_en']
            sch['school_name'] = std_school_names[c]['name_hi']
        
        # Standardize village
        v_raw = str(sch.get('village', '')).strip()
        clean_key = re.sub(r'[^A-Z0-9]', '', v_raw.upper())
        
        std_hi = None
        if v_raw in VILLAGE_ALIASES:
            std_hi = VILLAGE_ALIASES[v_raw]
        elif clean_key in VILLAGE_ALIASES:
            std_hi = VILLAGE_ALIASES[clean_key]
        elif v_raw.upper() in VILLAGE_ALIASES:
            std_hi = VILLAGE_ALIASES[v_raw.upper()]
        elif clean_key in census_by_norm_en:
            std_hi = census_by_norm_en[clean_key]['v_hi']
        elif v_raw in census_by_norm_hi:
            std_hi = v_raw
        
        if std_hi:
            sch['village_hi'] = std_hi
            sch['village'] = std_hi
            # Find en
            for cv in census_97:
                if cv['v_hi'] == std_hi:
                    sch['village_en'] = cv['v_en']
                    break
            total_updated_villages += 1

print(f"Total school village entries standardized: {total_updated_villages}")

# Save updated master_cbeo_data.json and master_cbeo_data.js
with open('master_cbeo_data.json', 'w', encoding='utf-8') as f:
    json.dump(master, f, ensure_ascii=False, indent=2)

with open('master_cbeo_data.js', 'w', encoding='utf-8') as f:
    f.write('const MASTER_CBEO_DATA = ' + json.dumps(master, ensure_ascii=False, indent=2) + ';\n')
    f.write('if (typeof module !== "undefined" && module.exports) { module.exports = MASTER_CBEO_DATA; }\n')

print("✓ master_cbeo_data.json and master_cbeo_data.js successfully updated with Standard 97 Hindi Village Names and Standard School Names!")
