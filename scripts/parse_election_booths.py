import openpyxl, json

# Mapping of Kruti Dev to Unicode for known terms
kruti_to_unicode_map = {
    "cM+xkao ¼lqj[k.M½": "बड़गांव (सुरखण्ड)",
    "cMxkao ¼lqj[k.M½": "बड़गांव (सुरखण्ड)",
    "cM+yh": "बड़ली",
    "cMyh": "बड़ली",
    "cxjkbZ": "बगराई",
    "ckUnuokMk": "बांदनवाड़ा",
    "fHkuk;": "भिनाय",
    "cwcfd;k": "बूबकिया",
    "pkikusjh": "चांपानेरी",
    "pkaikusjh": "चांपानेरी",
    "NNqUnjk": "छछुन्दरा",
    "nsofy;kdyka": "देवलियाकलां",
    "nsofy;kadyka": "देवलियाकलां",
    "nsoiqjk": "देवपुरा",
    "/kkarksy": "धांतोल",
    ",dyflgk": "एकलसिंगा",
    ",dyflagk": "एकलसिंगा",
    "?k.kk": "घणा",
    "xq<k[kqnZ": "गुढ़ाखुर्द",
    "fg;kfy;k": "हियालिया",
    "dusbZdyka": "कनईकलां",
    "djkaVh": "करांटी",
    "dSjksV": "कैरोट",
    "[ksMh": "खेड़ी",
    "[ksMh ": "खेड़ी",
    "dqEgkfj;ka": "कुम्हारिया",
    "ykexjk": "लामगरा",
    "ukxksyk": "नागोला",
    "ukUnlh": "नान्दसी",
    "iaMkxk": "पंडागा",
    "ikMfy;k": "पाडलिया",
    "ikM+fy;k": "पाडलिया",
    "jkeekfy;k": "राममालिया",
    "jkrkdksV": "राताकोट",
    "flaxkoy": "सिंगावल",
    "lkscMh": "सोबड़ी",
    "lkscM+h": "सोबड़ी",
    "lksy[kqnZ": "सोलखुर्द",
    "lksy[kqnZ ": "सोलखुर्द"
}

gp_to_school_code = {
    "बड़गांव (सुरखण्ड)": "221764",
    "बड़ली": "221755",
    "बगराई": "485030",
    "बांदनवाड़ा": "221769",
    "भिनाय": "221780",
    "बूबकिया": "221763",
    "चांपानेरी": "221758",
    "छछुन्दरा": "221787",
    "देवलियाकलां": "221754",
    "देवपुरा": "488941",
    "धांतोल": "221783",
    "एकलसिंगा": "221786",
    "घणा": "488897",
    "गुढ़ाखुर्द": "221762",
    "हियालिया": "488947",
    "कनईकलां": "221765",
    "करांटी": "221773",
    "कैरोट": "221767",
    "खेड़ी": "221774",
    "कुम्हारिया": "221777",
    "लामगरा": "221759",
    "नागोला": "221772",
    "नान्दसी": "221756",
    "पंडागा": "221788",
    "पाडलिया": "221766",
    "राममालिया": "221785",
    "राताकोट": "221775",
    "सिंगावल": "221781",
    "सोबड़ी": "221782",
    "सोलखुर्द": "410859"
}

def convert_kruti_to_hindi(text):
    if not text:
        return ""
    t = str(text).strip()
    for k, v in kruti_to_unicode_map.items():
        t = t.replace(k, v)
    
    rep = {
        "jktdh;": "राजकीय",
        "mPp": "उच्च",
        "ek/;fed": "माध्यमिक",
        "izkFkfed": "प्राथमिक",
        "fo|ky;": "विद्यालय",
        "egkRek": "महात्मा",
        "xka/kh": "गांधी",
        "ih,eJh": "पीएम श्री",
        "dejk": "कमरा",
        "u&": "नं.",
        "u0": "नं.",
        "¼": "(",
        "½": ")",
        "]": ",",
        "ljdkjh": "सरकारी",
        "vPNh": "अच्छी",
        "gka": "हाँ",
        "ugh": "नहीं",
        "Bhd": "ठीक स्थिति",
        "ehVj": "मीटर",
        "cMxkao": "बड़गांव",
        "cM+xkao": "बड़गांव",
        "cM+yh": "बड़ली",
        "cxjkbZ": "बगराई",
        "ckUnuokM+k": "बांदनवाड़ा",
        "ckUnuokMk": "बांदनवाड़ा",
        "fHkuk;": "भिनाय",
        "cwcfd;k": "बूबकिया",
        "pkikusjh": "चांपानेरी",
        "NNqUnjk": "छछुन्दरा",
        "nsofy;kdyka": "देवलियाकलां",
        "nsofy;kdyk": "देवलियाकलां",
        "nsoiqjk": "देवपुरा",
        "nsofj;k": "देवरिया",
        "/kkarksy": "धांतोल",
        ",dyflagk": "एकलसिंगा",
        ",dyflgk": "एकलसिंगा",
        "?k.kk": "घणा",
        "xq<k[kqnZ": "गुढ़ाखुर्द",
        "fg;kfy;k": "हियालिया",
        "dusbZdyka": "कनईकलां",
        "djkaVh": "करांटी",
        "dSjksV": "कैरोट",
        "[ksM+h": "खेड़ी",
        "[ksMh": "खेड़ी",
        "dqEgkfj;k": "कुम्हारिया",
        "dqEgkfj;ka": "कुम्हारिया",
        "ykexjk": "लामगरा",
        "ukxksyk": "नागोला",
        "ukUnlh": "नान्दसी",
        "iM+kxk": "पंडागा",
        "iaMkxk": "पंडागा",
        "ikMfy;k": "पाडलिया",
        "jkeekfy;k": "राममालिया",
        "jkrkdksV": "राताकोट",
        "flaxkoy": "सिंगावल",
        "lkscMh": "सोबड़ी",
        "lksy[kqnZ": "सोलखुर्द"
    }
    for k, v in rep.items():
        t = t.replace(k, v)
    return " ".join(t.split())

wb = openpyxl.load_workbook(r'C:\Users\jiten\Desktop\cbeo\matdan kendra virfy\PS DETAILS AJMER (1) 116 booth.xlsx', data_only=True)
s_p1 = wb['P-1']
s_p3 = wb['P-3']

# Build P-1 ward map
p1_wards = {}
for r in range(6, s_p1.max_row+1):
    bno = s_p1.cell(r, 4).value
    ward = s_p1.cell(r, 6).value
    if bno is not None:
        w_clean = str(ward).replace("]", ", ").strip() if ward else ""
        if w_clean.endswith(","): w_clean = w_clean[:-1].strip()
        p1_wards[int(bno)] = w_clean

booth_records = []
cur_gp_hi = ""
cur_bldg_hi = ""

for r in range(7, s_p3.max_row+1):
    gp_cell = s_p3.cell(r, 2).value
    room_cell = s_p3.cell(r, 3).value
    bldg_type_cell = s_p3.cell(r, 4).value
    booth_no = s_p3.cell(r, 5).value
    bldg_cond_cell = s_p3.cell(r, 7).value
    area_cell = s_p3.cell(r, 8).value
    road_dist_cell = s_p3.cell(r, 10).value
    boundary_cell = s_p3.cell(r, 11).value
    water_cell = s_p3.cell(r, 12).value
    toilet_cell = s_p3.cell(r, 13).value
    shed_cell = s_p3.cell(r, 14).value
    ramp_cell = s_p3.cell(r, 15).value
    light_bldg_cell = s_p3.cell(r, 16).value
    air_light_cell = s_p3.cell(r, 17).value
    furniture_cell = s_p3.cell(r, 18).value
    doors_cell = s_p3.cell(r, 19).value
    mf_line_cell = s_p3.cell(r, 20).value
    fan_cell = s_p3.cell(r, 21).value
    win_door_cell = s_p3.cell(r, 22).value
    remarks_cell = s_p3.cell(r, 23).value
    ps_cell = s_p3.cell(r, 24).value
    zp_cell = s_p3.cell(r, 25).value
    ac_cell = s_p3.cell(r, 26).value or 104

    if gp_cell:
        cur_gp_hi = convert_kruti_to_hindi(str(gp_cell).strip())
    
    if booth_no is not None:
        b_no = int(booth_no)
        room_hi = convert_kruti_to_hindi(str(room_cell).strip()) if room_cell else f"कक्ष नं. {b_no}"
        
        sch_code = gp_to_school_code.get(cur_gp_hi, "")
        if "महात्मा गांधी" in room_hi:
            if "बांदनवाड़ा" in room_hi: sch_code = "221769"
            elif "भिनाय" in room_hi: sch_code = "221780"
            elif "देवलियाकलां" in room_hi: sch_code = "221754"

        booth_records.append({
            "booth_no": b_no,
            "panchayat_hi": cur_gp_hi,
            "school_code": sch_code,
            "room_hi": room_hi,
            "ward": p1_wards.get(b_no, ""),
            "bldg_type": convert_kruti_to_hindi(str(bldg_type_cell or 'सरकारी')),
            "bldg_condition": convert_kruti_to_hindi(str(bldg_cond_cell or 'अच्छी')),
            "area_sqm": int(area_cell) if area_cell and str(area_cell).isdigit() else 50,
            "road_dist": convert_kruti_to_hindi(str(road_dist_cell or '100 मीटर')),
            "boundary": (str(boundary_cell).lower() != 'ugh'),
            "water": (str(water_cell).lower() != 'ugh'),
            "toilet": (str(toilet_cell).lower() != 'ugh'),
            "shed": (str(shed_cell).lower() != 'ugh'),
            "ramp": (str(ramp_cell).lower() != 'ugh'),
            "light_bldg": (str(light_bldg_cell).lower() != 'ugh'),
            "air_light": (str(air_light_cell).lower() != 'ugh'),
            "furniture": (str(furniture_cell).lower() != 'ugh'),
            "doors": (str(doors_cell).lower() != 'ugh'),
            "mf_line": (str(mf_line_cell).lower() != 'ugh'),
            "fan": (str(fan_cell).lower() != 'ugh'),
            "win_door": "ठीक स्थिति" if str(win_door_cell).lower() != 'ugh' else "मरम्मत योग्य",
            "ps_constituency": int(ps_cell) if ps_cell else "",
            "zp_constituency": int(zp_cell) if zp_cell else "",
            "ac_constituency": int(ac_cell) if ac_cell else 104
        })

print(f"Total processed booths from P-3: {len(booth_records)}")
with open("matdan_kendra_116_booths.json", "w", encoding="utf-8") as f:
    json.dump(booth_records, f, ensure_ascii=False, indent=2)

print("Saved matdan_kendra_116_booths.json successfully!")
