# -*- coding: utf-8 -*-
import openpyxl
import json
import re

wb = openpyxl.load_workbook('school enrolment.xlsx')
s = wb['समान परीक्षा मिलान व अंतर']

excel_schools = {}
for r in range(7, 56):
    code = str(s.cell(row=r, column=2).value or '').strip()
    raw = str(s.cell(row=r, column=3).value or '').strip()
    if code and code.isdigit() and raw:
        # Match trailing English name in parentheses
        m_en = re.search(r'\((GOVT\.[^)]*|MAHATMA[^)]*|PM\s*SHRI[^)]*)\)\s*$', raw, re.IGNORECASE)
        if m_en:
            en = m_en.group(1).strip()
            hi = raw[:m_en.start()].strip()
        else:
            hi = raw
            en = ''
        
        # Add comma after विद्यालय if not present, e.g. "राजकीय उच्च माध्यमिक विद्यालय, बड़गांव (सूरखण्ड)"
        hi_expanded = hi.replace('रा.बा.उ.मा.वि.', 'राजकीय बालिका उच्च माध्यमिक विद्यालय,') \
                        .replace('रा.उ.मा.वि.', 'राजकीय उच्च माध्यमिक विद्यालय,') \
                        .replace('रा.बा.मा.वि.', 'राजकीय बालिका माध्यमिक विद्यालय,') \
                        .replace('रा.मा.वि.', 'राजकीय माध्यमिक विद्यालय,') \
                        .replace('महात्मा गांधी रा.वि.', 'महात्मा गांधी राजकीय विद्यालय,') \
                        .replace('पीएम श्री रा.उ.मा.वि.', 'पीएम श्री राजकीय उच्च माध्यमिक विद्यालय,')
        # Clean double commas or spaces
        hi_expanded = re.sub(r',\s*,', ',', hi_expanded)
        hi_expanded = re.sub(r'विद्यालय,\s*', 'विद्यालय, ', hi_expanded).strip()
        
        # Format: कार्यालय <नाम> | शा.दा. कोड: <code>
        formal = f"कार्यालय {hi_expanded} | शा.दा. कोड: {code}"
        excel_schools[code] = {
            'code': code,
            'raw_hi': hi,
            'hi_expanded': hi_expanded,
            'en': en,
            'formal': formal
        }

print(f"Loaded {len(excel_schools)} official schools.")

# Update master_cbeo_data.json
with open('master_cbeo_data.json', encoding='utf-8') as f:
    master = json.load(f)

for sch in master.get('schools_56', []):
    c = str(sch.get('shala_darpan_code', '')).strip()
    if c in excel_schools:
        info = excel_schools[c]
        sch['school_name_hi'] = info['hi_expanded']
        sch['school_name'] = info['hi_expanded']
        sch['school_name_formal'] = info['formal']
        if info['en']:
            sch['school_name_en'] = info['en']

for p in master.get('peeos', []):
    for sch in p.get('schools', []):
        c = str(sch.get('shala_darpan_code', '')).strip()
        if c in excel_schools:
            info = excel_schools[c]
            sch['school_name_hi'] = info['hi_expanded']
            sch['school_name'] = info['hi_expanded']
            sch['school_name_formal'] = info['formal']
            if info['en']:
                sch['school_name_en'] = info['en']

with open('master_cbeo_data.json', 'w', encoding='utf-8') as f:
    json.dump(master, f, ensure_ascii=False, indent=2)

with open('master_cbeo_data.js', 'w', encoding='utf-8') as f:
    f.write('const MASTER_CBEO_DATA = ' + json.dumps(master, ensure_ascii=False, indent=2) + ';\n')
    f.write('if (typeof module !== "undefined" && module.exports) { module.exports = MASTER_CBEO_DATA; }\n')

print("✓ master_cbeo_data.json and master_cbeo_data.js updated!")
print("Sample formal 221764:", excel_schools.get('221764', {}).get('formal'))
print("Sample formal 221769:", excel_schools.get('221769', {}).get('formal'))
print("Sample formal 488781:", excel_schools.get('488781', {}).get('formal'))
